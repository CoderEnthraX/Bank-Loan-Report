"""Workbook sheets built from the strategy catalogue (catalog_<TF>.pkl): Strategy guide, <TF> strategies (every indicator input, exactly as named
in the indicator, per scenario + what to expect), All profitable settings. Called from build_final.py."""
import os, sys, pickle, collections
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter as CL
S = '/tmp/claude-0/-home-user-Bank-Loan-Report/21ed2550-acec-5cf1-bf01-a24cbf22fe18/scratchpad'
sys.path.insert(0, S + '/gold')
import greedy, labels, pickcfg
from build_final import put, hdr, title, font, widths, BOX, HDR, SUB, TOT, GOOD, BAD, BLUE, GREEN, F, DAYS_N, TFS
CHG = PatternFill('solid', fgColor='FFF2CC'); GREY = '808080'
def tick(v): return 'On' if v else 'Off'
def num(v): return labels.fmt(float(v)) if isinstance(v, float) else str(v)
ENTRY = {'model': 'Model (body mid or 6-bar level)', 'pct': 'Custom % of signal candle', 'market': 'Market (next bar open)'}
KZN = ('Asia 05:30-09:30', 'London 11:30-14:30', 'New York 16:30-19:30', 'London Close 19:30-21:30')
EAKZ = ('InpKzAsia', 'InpKzLon', 'InpKzNy', 'InpKzLc')
# (group, input exactly as in the indicator, MT5 EA input, value shown, research keys it covers)
ROWS = [('SIGNAL MODEL', 'Structure Lookback (bars)', 'InpStructLen', lambda d: d['structLen'], ('structLen',)),
        ('SIGNAL MODEL', 'Signal Break Lookback (bars)', 'InpBrkLen', lambda d: d['brkLen'], ('brkLen',)),
        ('SIGNAL MODEL', 'Entry Level Lookback (bars)', 'InpEntryLen', lambda d: d['entryLen'], ('entryLen',)),
        ('SIGNAL MODEL', 'Retest Window (bars)', 'InpRetestBars', lambda d: d['retestBars'], ('retestBars',)),
        ('SIGNAL MODEL', 'Min Displacement Range (x ATR)', 'InpDispAtr', lambda d: d['dispAtr'], ('dispAtr',)),
        ('SIGNAL MODEL', 'Min Displacement Body (% of range)', 'InpBodyPct', lambda d: d['dispBodyPct'], ('dispBodyPct',)),
        ('SIGNAL MODEL', 'Max Close Distance From Extreme (% of range)', 'InpClosePct', lambda d: d['dispClosePct'], ('dispClosePct',)),
        ('SIGNAL MODEL', 'Momentum Lookback (bars) · Min Move (x ATR)   [1st box]', 'InpMomLen', lambda d: d['momLen'], ('momLen',)),
        ('SIGNAL MODEL', '   ... 2nd box: Min Move (x ATR)   [same line]', 'InpMomAtr', lambda d: d['momAtr'], ('momAtr',)),
        ('SIGNAL MODEL', 'Min CMO (clean move)', 'InpCmoMin', lambda d: d['cmoMin'], ('cmoMin',)),
        ('SIGNAL MODEL', 'Same-Direction Re-arm', 'InpRearm', lambda d: 'Keep first' if d['rearmKeep'] else 'Replace with newer', ('rearmKeep',)),
        ('SIGNAL MODEL', 'Stop Mode', 'InpStopMode', lambda d: 'ATR Model' if d['stopMode'] == 'atr' else 'Swing + Buffer (ATR capped)', ('stopMode',)),
        ('SIGNAL MODEL', 'Stop Distance (x ATR)', 'InpStopAtr', lambda d: d['stopAtr'], ('stopAtr',)),
        ('SIGNAL MODEL', 'Swing Lookback (bars) · Buffer (x ATR)   [1st box]', 'InpStopLen', lambda d: d['stopLen'], ('stopLen',)),
        ('SIGNAL MODEL', '   ... 2nd box: Buffer (x ATR)   [same line]', 'InpStopBufAtr', lambda d: d['stopBufAtr'], ('stopBufAtr',)),
        ('SIGNAL MODEL', 'Half Stop On First Counter-Bias Impulse', 'InpHalfStop', lambda d: tick(d['halfStopOn']), ('halfStopOn',)),
        ('SIGNAL MODEL', 'Half Stop (x ATR) · Prev Session Net Impulses   [1st box]', 'InpHalfStopAtr', lambda d: d['halfStopAtr'], ('halfStopAtr',)),
        ('SIGNAL MODEL', '   ... 2nd box: Prev Session Net Impulses   [same line]', 'InpBiasNet', lambda d: d['biasNet'], ('biasNet',)),
        ('SIGNAL MODEL', 'Impulse Min Body (% of range)', 'InpImpBody', lambda d: d['impBodyPct'], ('impBodyPct',)),
        ('SIGNAL MODEL', 'Half Stop When Session Opens Outside Prev Range', 'InpHalfOpenOut', lambda d: tick(d['halfOpenOut']), ('halfOpenOut',)),
        ('SIGNAL MODEL', 'Strict Bias Check (optional)', 'InpHalfStrict', lambda d: tick(d['halfStrict']), ('halfStrict',)),
        ('SIGNAL MODEL', 'Half Stop Never Wider Than The Normal Stop (optional)', 'InpHalfCap', lambda d: tick(d['halfCap']), ('halfCap',)),
        ('SIGNAL MODEL', 'Cooldown After Exit (bars)', 'InpCooldown', lambda d: d['cooldownBars'], ('cooldownBars',)),
        ('SIGNAL MODEL', 'Max Fills Per Day', 'InpMaxFills', lambda d: d['maxFills'], ('maxFills',)),
        ('SIGNAL MODEL', 'Daily Stop (R lost, 0 = off)', 'InpDayStopR', lambda d: d['dayStopR'], ('dayStopR',)),
        ('SIGNAL MODEL', 'Count Fill-Bar Stops In Daily Stop', 'InpFillBarStopR', lambda d: tick(d['fillBarStopR']), ('fillBarStopR',)),
        ('SIGNAL MODEL', 'Re-anchor Breakout After A Trade', 'InpLegReset', lambda d: tick(d['legReset']), ('legReset',)),
        ('SIGNAL MODEL', 'ATR Length', 'InpAtrLen', lambda d: d['atrLen'], ('atrLen',)),
        ('SIGNAL MODEL', 'Longs Above VWAP / Shorts Below', 'InpVwapFilter', lambda d: tick(d['vwapFilter']), ('vwapFilter',)),
        ('SIGNAL MODEL', 'VWAP Slope Bars', 'InpSlopeBars', lambda d: d['slopeBars'], ('slopeBars',)),
        ('SIGNAL MODEL', 'Slope Looks Across The Session Start', 'InpSlopeCross', lambda d: tick(d['slopeCross']), ('slopeCross',)),
        ('SIGNAL MODEL', 'Skip First Bars After VWAP Reset', 'InpMinOpenBars', lambda d: d['minOpenBars'], ('minOpenBars',)),
        ('SIGNAL MODEL', 'Skip Sunday (UTC) Weekly Open', 'InpSkipSunday', lambda d: tick(d['skipSunday']), ('skipSunday',)),
        ('SIGNAL MODEL', 'Deep Pullback Signal (optional) · Min Depth (x ATR)   [tick box]', 'InpDeepOn', lambda d: tick(d['deepOn']), ('deepOn',)),
        ('SIGNAL MODEL', '   ... box: Min Depth (x ATR)   [same line]', 'InpDeepFar', lambda d: d['deepFar'], ('deepFar',)),
        ('SIGNAL MODEL', 'Longs', 'InpAllowLong', lambda d: tick(d['allowLong']), ('allowLong',)),
        ('SIGNAL MODEL', 'Shorts', 'InpAllowShort', lambda d: tick(d['allowShort']), ('allowShort',)),
        ('VWAP', 'Anchor', 'InpVwapAnchor', lambda d: d['vwAnchorG'].capitalize(), ('vwAnchorG',)),
        ('VWAP', 'Source', 'InpVwapSrc', lambda d: d['vwSrcG'], ('vwSrcG',)),
        ('ENTRY PRICE', 'Entry Price', 'InpEntryMode', lambda d: ENTRY[d['entryMode']], ('entryMode',)),
        ('ENTRY PRICE', 'Custom % (0 = close, 50 = mid, 100 = low/high)', 'InpEntryPct', lambda d: d['entryPct'], ('entryPct',)),
        ('TARGETS, EXIT AND STOP MOVES', 'Targets TP1 · TP2 · TP3 · TP4 (R): 1st box (TP1)', 'InpTp1R', lambda d: d['tp1R'], ('tp1R',)),
        ('TARGETS, EXIT AND STOP MOVES', '   ... 2nd box (TP2)   [same line]', 'InpTp2R', lambda d: d['tp2R'], ('tp2R',)),
        ('TARGETS, EXIT AND STOP MOVES', '   ... 3rd box (TP3)   [same line]', 'InpTp3R', lambda d: d['tp3R'], ('tp3R',)),
        ('TARGETS, EXIT AND STOP MOVES', '   ... 4th box (TP4)   [same line]', 'InpTp4R', lambda d: d['tp4R'], ('tp4R',)),
        ('TARGETS, EXIT AND STOP MOVES', 'Final Exit At', 'InpExitAt', lambda d: 'TP%d' % d['exitAt'], ('exitAt',)),
        ('TARGETS, EXIT AND STOP MOVES', 'When TP1 Hits, Move Stop To (R)   [tick box]', 'InpMv1On', lambda d: tick(d['mv1On']), ('mv1On',)),
        ('TARGETS, EXIT AND STOP MOVES', '   ... value (R)   [same line]', 'InpMv1R', lambda d: d['mv1R'], ('mv1R',)),
        ('TARGETS, EXIT AND STOP MOVES', 'When TP2 Hits, Move Stop To (R)   [tick box]', 'InpMv2On', lambda d: tick(d['mv2On']), ('mv2On',)),
        ('TARGETS, EXIT AND STOP MOVES', '   ... value (R)   [same line]', 'InpMv2R', lambda d: d['mv2R'], ('mv2R',)),
        ('TARGETS, EXIT AND STOP MOVES', 'When TP3 Hits, Move Stop To (R)   [tick box]', 'InpMv3On', lambda d: tick(d['mv3On']), ('mv3On',)),
        ('TARGETS, EXIT AND STOP MOVES', '   ... value (R)   [same line]', 'InpMv3R', lambda d: d['mv3R'], ('mv3R',)),
        ('ACCURACY FILTERS', "Don't Chase: Skip Signals Beyond (σ From VWAP)   [tick box]", 'InpNoChase', lambda d: tick(d['noChase']), ('noChase',)),
        ('ACCURACY FILTERS', '   ... box: σ From VWAP   [same line]', 'InpNoChaseSig', lambda d: d['noChaseSig'], ('noChaseSig',)),
        ('ACCURACY FILTERS', 'After A Loss, No Same-Direction Trade Until A Win', 'InpLossBlock', lambda d: tick(d['lossBlock']), ('lossBlock',)),
        ('ACCURACY FILTERS', 'Skip Signals Too Close To VWAP: Min Distance (x ATR)   [tick box]', 'InpMinVwOn', lambda d: tick(d['minVwOn']), ('minVwOn',)),
        ('ACCURACY FILTERS', '   ... box: x ATR   [same line]', 'InpMinVwAtr', lambda d: d['minVwAtr'], ('minVwAtr',)),
        ('ACCURACY FILTERS', 'Skip Choppy Sessions: Max VWAP Crosses (Per 100 Bars)   [tick box]', 'InpChopOn', lambda d: tick(d['chopOn']), ('chopOn',)),
        ('ACCURACY FILTERS', '   ... box: Per 100 Bars   [same line]', 'InpChopMax', lambda d: d['chopMax'], ('chopMax',)),
        ('ACCURACY FILTERS', 'Skip When The 1h EMA20 Moved More Than (ATR In 3 Hours)   [tick box]', 'InpEmaOn', lambda d: tick(d['emaOn']), ('emaOn',)),
        ('ACCURACY FILTERS', '   ... box: ATR In 3 Hours   [same line]', 'InpEmaMax', lambda d: d['emaMax'], ('emaMax',)),
        ('DAYS AND HOURS', 'Time Zone', 'InpUserGMT = 5.5', lambda d: 'Asia/Kolkata', ()),
        ('DAYS AND HOURS', 'Trade Only On The Ticked Days', 'InpDayOn', lambda d: tick(d['dayOn']), ('dayOn',)),
        ('DAYS AND HOURS', '   Mon, Tue, Wed, Thu, Fri, Sat, Sun   [one tick box each]', 'InpMon ... InpSun',
         lambda d: 'all ticked' if all(d['days']) else ' '.join(n for n, v in zip(('Mon', 'Tue', 'Wed', 'Thu', 'Fri', 'Sat', 'Sun'), d['days']) if v), ('days',)),
        ('DAYS AND HOURS', 'Trade Only Between   [tick box]', 'InpHoursOn', lambda d: tick(d['hoursOn']), ('hoursOn',)),
        ('DAYS AND HOURS', '   session, Indian time   [same line]', 'InpStartH:M - InpEndH:M', lambda d: '%04d-%04d' % tuple(d['hoursSess']), ('hoursSess',)),
        ('DAYS AND HOURS', 'Trade Only In Kill Zones (Indian time)', 'InpKzOn', lambda d: tick(d['kzOn']), ('kzOn',))] + \
       [('DAYS AND HOURS', '   ' + KZN[k], EAKZ[k], (lambda k: lambda d: tick(d['kz'][k]))(k), ('kz',)) for k in range(4)] + \
       [('DAYS AND HOURS', 'Shift All Kill Zones (hours)', 'InpKzShift', lambda d: d['kzShift'], ('kzShift',)),
        ('PARTIAL PROFITS (strategy / EA only)', 'Close % At TP1', 'InpP1', lambda d: int(round(100 * d['p1'])) if d['exitAt'] > 1 else 0, ('p1',)),
        ('PARTIAL PROFITS (strategy / EA only)', 'Close % At TP2', 'InpP2', lambda d: 0, ()),
        ('PARTIAL PROFITS (strategy / EA only)', 'Close % At TP3', 'InpP3', lambda d: 0, ())]
COVERED = {k for r in ROWS for k in r[4]}
def full(cfg):
    d = dict(greedy.defaults()); d.update(cfg); return d
def check_covered(cfg):
    d0 = greedy.defaults(); miss = [k for k, v in cfg.items() if d0.get(k) != v and k not in COVERED]
    assert not miss, ('settings not shown in the sheet', miss)
def show(v): return num(v) if isinstance(v, (int, float)) and not isinstance(v, bool) else v
def load():
    out = {}
    for tf in TFS:
        fn = S + '/gold/catalog_%s.pkl' % tf
        if os.path.exists(fn): out[tf] = pickle.load(open(fn, 'rb'))
    return out
def sid(tf, k): return '%s-S%d' % (tf, k + 1)
def works(cfg): return 'Strategy and EA only (partial close)' if full(cfg)['p1'] and full(cfg)['exitAt'] > 1 else 'Indicator, strategy and EA'
def counts(tr):
    net = [x[1] - x[4] / x[5] for x in tr]
    A = [v for x, v in zip(tr, net) if x[0].year <= 2023]; B = [v for x, v in zip(tr, net) if x[0].year >= 2024]
    return dict(n=len(tr), w=sum(1 for x in tr if x[2]), p=sum(1 for x in tr if x[1] > 0), h=[sum(1 for x in tr if x[3][k]) for k in range(4)],
                be=sum(1 for x in tr if x[6]), net=sum(net), nA=len(A), netA=sum(A), nB=len(B), netB=sum(B))
# ---------------------------------------------------------------- one sheet per timeframe
def tf_sheet(wb, tf, C):
    ws = wb.create_sheet('%s strategies' % tf); st = C['strategies']; k0 = 5          # strategies start in column E
    title(ws, '%s: profitable settings for each scenario - what to set in the indicator and what to expect' % tf)
    ws['A2'] = ('Pick the column for your goal and read down. Yellow cells = change this input to the value shown; grey = leave the default. All results are 2021-2026 on FxPro '
                'gold after the spread, and every setting here was profitable in both 2021-23 and 2024-26 with at least 4 of 6 years positive. Expect live results to be weaker.')
    ws['A2'].font = font(italic=True); ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=k0 + len(st)); ws['A2'].alignment = Alignment(wrap_text=True, vertical='top')
    ws.row_dimensions[2].height = 30
    r = 4; hdr(ws, r, ['Group', 'Input (exactly as in the indicator)', 'MT5 EA input', 'Default'] + [sid(tf, k) for k in range(len(st))], height=20)
    head = r; r += 1
    def section(name):
        nonlocal r
        put(ws, r, 1, name, bold=True, fill=SUB); ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=k0 - 1 + len(st)); r += 1
    section('WHEN TO USE IT')
    for lab, fn in (('Scenario', lambda s: ' + '.join(n for n, _ in s['scen'])), ('Use it when', lambda s: ' '.join(w for _, w in s['scen'])), ('Works in', lambda s: works(s['cfg']))):
        put(ws, r, 2, lab, bold=True, wrap=True)
        for k, s in enumerate(st): put(ws, r, k0 + k, fn(s), bold=(lab == 'Scenario'), wrap=True)
        ws.row_dimensions[r].height = 64 if lab == 'Use it when' else 30 if lab == 'Scenario' else 15; r += 1
    section('SETTINGS - change the yellow cells, leave the rest at the default')
    d0 = full({}); grp = None
    for s in st: check_covered(s['cfg'])
    for g, lab, ea, fn, keys in ROWS:
        put(ws, r, 1, g if g != grp else '', bold=(g != grp)); grp = g
        put(ws, r, 2, lab); put(ws, r, 3, ea, color=GREY); dv = show(fn(d0)); put(ws, r, 4, dv, color=GREY, align='center')
        for k, s in enumerate(st):
            v = show(fn(full(s['cfg'])))
            if v == dv: put(ws, r, k0 + k, v, color=GREY, align='center')
            else: put(ws, r, k0 + k, v, bold=True, fill=CHG, align='center')
        r += 1
    section('WHAT TO EXPECT - 2021-2026, FxPro gold, after the spread (blue = backtest result, black = formula)')
    cs = [counts(s['trades']) for s in st]; first = r
    def row(lab, vals, nf=None, color=BLUE, bold=False):
        nonlocal r
        put(ws, r, 2, lab, bold=bold)
        for k, v in enumerate(vals):
            f = isinstance(v, str) and v.startswith('=')
            put(ws, r, k0 + k, v, nf, bold, None if f else color, align='center')
        r += 1; return r - 1
    L = lambda k: CL(k0 + k)
    rn = row('Trades', [c['n'] for c in cs], '#,##0')
    row('Trades per day', ['=%s%d/%d' % (L(k), rn, DAYS_N) for k in range(len(st))], '0.00')
    rw = row('Wins (closed at the final target)', [c['w'] for c in cs], '#,##0')
    row('Win % (closed at the final target)', ['=IF(%s%d=0,"",%s%d/%s%d)' % (L(k), rn, L(k), rw, L(k), rn) for k in range(len(st))], '0.0%')
    rp = row('Profitable trades', [c['p'] for c in cs], '#,##0')
    row('Profitable trades %', ['=IF(%s%d=0,"",%s%d/%s%d)' % (L(k), rn, L(k), rp, L(k), rn) for k in range(len(st))], '0.0%', bold=True)
    for j in range(4):
        rh = row('TP%d reached (count)' % (j + 1), [c['h'][j] for c in cs], '#,##0')
        row('TP%d reached %% (TP%d = R shown in TARGETS)' % (j + 1, j + 1), ['=IF(%s%d=0,"",%s%d/%s%d)' % (L(k), rn, L(k), rh, L(k), rn) for k in range(len(st))], '0.0%')
    rb = row('Breakeven exits (stop moved to entry)', [c['be'] for c in cs], '#,##0')
    rt = row('Total R after spread', [round(c['net'], 3) for c in cs], '0.0')
    re = row('R per trade after spread, 2021-2026', ['=IF(%s%d=0,"",%s%d/%s%d)' % (L(k), rn, L(k), rt, L(k), rn) for k in range(len(st))], '0.000', bold=True)
    ra_n = row('2021-23 trades', [c['nA'] for c in cs], '#,##0'); ra = row('2021-23 total R after spread', [round(c['netA'], 3) for c in cs], '0.0')
    row('R per trade, 2021-23', ['=IF(%s%d=0,"",%s%d/%s%d)' % (L(k), ra_n, L(k), ra, L(k), ra_n) for k in range(len(st))], '0.000')
    rb_n = row('2024-26 trades', [c['nB'] for c in cs], '#,##0'); rb2 = row('2024-26 total R after spread', [round(c['netB'], 3) for c in cs], '0.0')
    row('R per trade, 2024-26', ['=IF(%s%d=0,"",%s%d/%s%d)' % (L(k), rb_n, L(k), rb2, L(k), rb_n) for k in range(len(st))], '0.000')
    row('R per trade if live costs are 0.05 R higher', ['=IF(%s%d="","",%s%d-0.05)' % (L(k), re, L(k), re) for k in range(len(st))], '0.000')
    ys = sorted({y for s in st for y in s['stats']['years']}); yr0 = r
    for y in ys: row('%d: total R after spread' % y, [round(s['stats']['years'].get(y, 0.0), 2) for s in st], '0.0')
    row('Years positive (of %d)' % len(ys), ['=COUNTIF(%s%d:%s%d,">0")' % (L(k), yr0, L(k), r - 1) for k in range(len(st))], None, bold=True)
    row('Longest winning streak', [s['stats']['maxW'] for s in st], None, bold=True)
    row('Longest losing streak', [s['stats']['maxL'] for s in st], None, bold=True)
    row('Winning streaks of 5 or more (how often)', [s['stats']['W5'] for s in st])
    row('Losing streaks of 5 or more (how often)', [s['stats']['L5'] for s in st])
    row('Deepest fall from a peak (R, fixed risk)', [round(s['stats']['ddR'], 2) for s in st], '0.0')
    section('MONEY - $10,000 start, $50 risk per trade (strategy / EA sizing: a stop-out loses $50 including the spread)')
    m = lambda key, f: [f(s['money'][key]) for s in st]
    row('Fixed $50 risk (Loss Recovery = Off): net profit ($)', m('off', lambda x: round(x['net'], 2)), '#,##0;(#,##0);-')
    row('   deepest fall from a peak ($)', m('off', lambda x: round(x['maxDD'], 2)), '#,##0;(#,##0);-')
    row('   deepest fall from a peak (%)', m('off', lambda x: round(x['maxDDpct'], 4)), '0.0%')
    row('Rule C split, split 3, $500 hard cap, Start again: net profit ($)', m('cs3_reset', lambda x: 'BLOWN' if x['ruined'] else round(x['net'], 2)), '#,##0;(#,##0);-')
    row('   deepest fall from a peak (%)', m('cs3_reset', lambda x: round(x['maxDDpct'], 4)), '0.0%')
    row('Rule C split, split 3, $500 hard cap, Clamp: net profit ($)', m('cs3_clamp', lambda x: 'BLOWN' if x['ruined'] else round(x['net'], 2)), '#,##0;(#,##0);-')
    row('   deepest fall from a peak (%)', m('cs3_clamp', lambda x: round(x['maxDDpct'], 4)), '0.0%')
    mc = lambda key, f: [f(s['money']['mc'][key]) if s['money'].get('mc') else '-' for s in st]
    for key, name, mk in (('crA_reset', 'Crooks A, $500 hard cap, Start again', ('Crooks A', 3, 500.0, 'reset')), ('crB3_reset', 'Crooks B, split 3, $500 hard cap, Start again', ('Crooks B', 3, 500.0, 'reset'))):
        row(name + ': net profit ($)', m(key, lambda x: 'BLOWN' if x['ruined'] else round(x['net'], 2)), '#,##0;(#,##0);-')
        row('   deepest fall from a peak (%)', m(key, lambda x: round(x['maxDDpct'], 4)), '0.0%')
        row('   the same trades in 500 random orders: median net ($)', mc(mk, lambda q: round(q['med'], 2)), '#,##0;(#,##0);-')
        row('   ... 95% of the random orders fall less than (%)', mc(mk, lambda q: round(q['dd95'], 4)), '0.0%')
    def advice(s):
        e = s['stats']['E']; rs = s['money']['cs3_reset']
        if tf in ('M1', 'M5') or e < 0.08:
            q = (s['money'].get('mc') or {}).get(('Crooks A', 3, 500.0, 'reset'))
            return ('Fixed risk (Off). The edge is too thin for Rule C: 0.05 R of extra cost per trade can blow the account.' +
                    (' Crooks A with the $500 cap (Start again): 1 in 20 random orders of these trades falls more than %.0f%% (Crooks A and B test).' % (100 * q['dd95']) if q else ''))
        if rs['ruined'] or rs['maxDDpct'] > 0.40: return 'Fixed risk (Off). Rule C split falls too far here.'
        return ('Fixed risk (Off) for the smallest falls; Rule C split (split 3, $500 cap, Start again) for growth, if you accept the deeper falls shown above. '
                'Crooks A / B ($500 cap, Start again) sit in between: see the Crooks A and B test sheet.')
    put(ws, r, 2, 'Recommended money management', bold=True)
    for k, s in enumerate(st): put(ws, r, k0 + k, advice(s), wrap=True)
    ws.row_dimensions[r].height = 100; r += 1
    widths(ws, [26, 52, 22, 16] + [26] * len(st)); ws.freeze_panes = '%s%d' % (CL(k0), head + 1)
    return {k: '%s!%s%d' % ("'%s strategies'" % tf, L(k), head) for k in range(len(st))}
# ---------------------------------------------------------------- overview
SCEN = ['Most profit', 'Shortest losing streak', 'Shortest losing streak - indicator only', 'Longest winning streak', 'Highest win rate', 'Most trades a day',
        'Every year positive', 'Smoothest', 'Simplest', 'Asia session only', 'London session only', 'New York session only', 'London close session only']
def guide(wb, CAT, links):
    ws = wb.create_sheet('Strategy guide', 1)
    title(ws, 'Strategy guide: which settings to use for which goal (FxPro MT5 gold, 2021-2026, after the spread)')
    ws['A2'] = ('Part 1: for each goal and timeframe, the strategy to use (click it for the exact settings on the timeframe sheet). Part 2: every strategy with its key numbers. '
                'Every strategy is profitable after spread in 2021-23 and 2024-26 and positive in at least 4 of 6 years. "-" = no profitable setting for that goal.')
    ws['A2'].font = font(italic=True); ws.merge_cells('A2:P2'); ws['A2'].alignment = Alignment(wrap_text=True, vertical='top'); ws.row_dimensions[2].height = 30
    r = 4; put(ws, r, 1, 'PART 1 - pick by goal', bold=True, fill=SUB); ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=5); r += 1
    tfs = [tf for tf in TFS if tf in CAT]; hdr(ws, r, ['Goal'] + tfs, height=20); r += 1
    for sc in SCEN:
        put(ws, r, 1, sc, bold=True)
        for j, tf in enumerate(tfs):
            st = CAT[tf]['strategies']; hit = [k for k, s in enumerate(st) if sc in [n for n, _ in s['scen']]]
            if not hit: put(ws, r, 2 + j, '-', color=GREY, align='center'); continue
            k = hit[0]; x = st[k]['stats']
            c = put(ws, r, 2 + j, '%s: W/L streak %d/%d, %+.3f R, %.2f/day' % (sid(tf, k), x['maxW'], x['maxL'], x['E'], x['perday']), wrap=True)
            c.hyperlink = '#' + links[tf][k]; c.font = Font(name=F, size=9, color='0563C1', underline='single')
        r += 1
    r += 1; put(ws, r, 1, 'PART 2 - every strategy', bold=True, fill=SUB); ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=5); r += 1
    cols = ['Strategy', 'Timeframe', 'Goal(s)', 'Works in', 'Trades per day', 'Profitable trades %', 'TP1 reached %',
            'R per trade 2021-23', 'R per trade 2024-26', 'R per trade 2021-2026', 'Years positive', 'Longest winning streak', 'Longest losing streak',
            'Deepest fall (R)', 'Fixed $50: net ($)', 'Rule C split 3, $500 cap, Start again: net ($)', '... its deepest fall %',
            'Crooks A, $500 cap, Start again: net ($)', '... its deepest fall %', 'Settings to change (indicator names)']
    hdr(ws, r, cols, height=54); r += 1
    for tf in tfs:
        for k, s in enumerate(CAT[tf]['strategies']):
            x = s['stats']; rs = s['money']['cs3_reset']; ca = s['money']['crA_reset']
            vals = [sid(tf, k), tf, ' + '.join(n for n, _ in s['scen']), works(s['cfg']), round(x['perday'], 3), round(x['prof'], 4), round(x['tp'][0], 4),
                    round(x['EA'], 4), round(x['EB'], 4), round(x['E'], 4), x['ypos'], x['maxW'], x['maxL'], round(x['ddR'], 2), round(s['money']['off']['net'], 2),
                    'BLOWN' if rs['ruined'] else round(rs['net'], 2), round(rs['maxDDpct'], 4), 'BLOWN' if ca['ruined'] else round(ca['net'], 2), round(ca['maxDDpct'], 4),
                    labels.short(s['cfg'])]
            for j, v in enumerate(vals, 1):
                c = put(ws, r, j, v, '0.00' if j == 5 else '0.0%' if j in (6, 7, 17, 19) else '0.000' if j in (8, 9, 10) else '0.0' if j == 14 else '#,##0;(#,##0);-' if j in (15, 16, 18) else None,
                        j == 1, BLUE if 5 <= j <= 19 else None, wrap=(j in (3, 4, 20)))
                if j == 1: c.hyperlink = '#' + links[tf][k]; c.font = Font(name=F, size=9, bold=True, color='0563C1', underline='single')
            ws.row_dimensions[r].height = 40; r += 1
        r += 1
    widths(ws, [34, 26, 26, 26, 26, 9, 8, 9, 9, 9, 8, 9, 9, 8, 10, 12, 9, 12, 9, 90]); ws.freeze_panes = 'B6'
# ---------------------------------------------------------------- every profitable setting
def allprof(wb, CAT):
    ws = wb.create_sheet('All profitable settings')
    title(ws, 'Every profitable setting found (2021-2026, after spread): profitable in 2021-23 AND 2024-26, at least 4 of 6 years positive, enough trades',
          'Use the filter arrows: e.g. Longest losing streak <= 6, or Works in = Indicator. Identical results are listed once (fewest changes). Ranked per timeframe by the weaker half.')
    r = 4
    cols = ['Timeframe', 'Rank', 'Settings to change (indicator names)', 'Number of changes', 'Works in', 'Trades', 'Trades per day', 'Profitable trades %', 'TP1 reached %',
            'R per trade 2021-23', 'R per trade 2024-26', 'R per trade 2021-2026', 'Total R after spread', 'Years positive', 'Longest winning streak', 'Longest losing streak']
    hdr(ws, r, cols, height=42); top = r; r += 1
    for tf in TFS:
        if tf not in CAT: continue
        rows = sorted(CAT[tf]['unique'], key=lambda q: -pickcfg.robust(q[1]))
        for k, (c, m) in enumerate(rows, 1):
            a = m['all']
            vals = [tf, k, labels.short(c), len(labels.rows(c)), works(c), a['n'], '=F%d/%d' % (r, DAYS_N), round(a['win'], 4), round(m['tp1'], 4), round(m['A']['E'], 4),
                    round(m['B']['E'], 4), round(a['E'], 4), round(a['net'], 2), m['ypos'], a['maxW'], a['maxL']]
            for j, v in enumerate(vals, 1):
                put(ws, r, j, v, '0.00' if j == 7 else '0.0%' if j in (8, 9) else '0.000' if j in (10, 11, 12) else '0.0' if j == 13 else None, False,
                    BLUE if (j >= 6 and j != 7) else None)
            r += 1
    ws.auto_filter.ref = 'A%d:%s%d' % (top, CL(len(cols)), r - 1)
    widths(ws, [9, 6, 110, 8, 30, 7, 8, 9, 8, 9, 9, 9, 9, 8, 9, 9]); ws.freeze_panes = 'D5'
