"""Workbook sheet 'VWAP distance' from vwfilter_scan.pkl / vwfilter_trades*.pkl: are trades far from the VWAP worse, and does a max-distance
cut-off (ATR, price, sigma) help in both 2021-23 and 2024-26? The findings text is built from the results. Called from build_final.py."""
import os, sys, pickle
from openpyxl.styles import Alignment
S = '/tmp/claude-0/-home-user-Bank-Loan-Report/21ed2550-acec-5cf1-bf01-a24cbf22fe18/scratchpad'
sys.path.insert(0, S + '/gold')
from build_final import put, hdr, title, font, widths, SUB, TOT, GOOD, BAD, BLUE, TFS
UNIT = {'atr': 'x ATR', 'price': 'price ($)', 'sigma': 'sigma (Don\'t Chase)'}
def note_rows(ws, r, lines, cols=16):
    for k, t in enumerate(lines):
        c = ws.cell(r + k, 1, t); c.font = font(bold=(k == 0), size=9); c.alignment = Alignment(wrap_text=True, vertical='top')
        ws.merge_cells(start_row=r + k, start_column=1, end_row=r + k, end_column=cols)
        if k: ws.row_dimensions[r + k].height = max(15, 12.5 * -(-len(t) // 190) + 3)
    return r + len(lines) + 1
def section(ws, r, text, cols=12):
    put(ws, r, 1, text, bold=True, fill=SUB); ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=cols); return r + 1
def vwfilter_sheet(wb):
    fs = S + '/gold/vwfilter_scan.pkl'
    if not os.path.exists(fs): return
    D = pickle.load(open(fs, 'rb')); T = pickle.load(open(S + '/gold/vwfilter_trades.pkl', 'rb'))
    if os.path.exists(S + '/gold/vwfilter_trades_tp2.pkl'): T.update(pickle.load(open(S + '/gold/vwfilter_trades_tp2.pkl', 'rb')))
    order = sorted(D, key=lambda k: (TFS.index(k[0]), 'exit TP2' in k[1], k[1] != 'Default settings', k[1]))
    prof = [k for k in order if D[k]['res']['none']['A']['E'] > 0 and D[k]['res']['none']['B']['E'] > 0]
    ws = wb.create_sheet('VWAP distance')
    title(ws, 'Do trades far from the VWAP lose more? Max-distance cut-off in ATR, in price and in sigma (FxPro MT5 gold, 2021-2026, after spread)')
    # summary across the profitable settings: all of them, and split by VWAP anchor (Session = the default; Month = most catalogue M1 / M5 / H1)
    anc = lambda k: T[k]['cfg'].get('vwAnchorG', 'session')
    GRP = [('All profitable settings', prof), ('Session anchor', [k for k in prof if anc(k) == 'session']), ('Month anchor', [k for k in prof if anc(k) == 'month'])]
    def summary(keys):
        out = {}
        for unit, grid in (('atr', (1.5, 2.0, 2.5, 3.0, 3.5, 4.0, 4.5, 5.0, 5.5, 6.0)), ('sigma', (1.5, 2.0, 2.5, 3.0, 3.5, 4.0))):
            for X in grid:
                better = worse = 0; R0 = R1 = n0 = n1 = 0.0
                for k in keys:
                    v = D[k]['res'].get((unit, X)); b = D[k]['res']['none']
                    if v is None: continue
                    better += v['A']['E'] > b['A']['E'] and v['B']['E'] > b['B']['E']; worse += v['A']['E'] < b['A']['E'] and v['B']['E'] < b['B']['E']
                    R0 += b['all']['R']; R1 += v['all']['R']; n0 += b['all']['n']; n1 += v['all']['n']
                out[(unit, X)] = (better, worse, R0, R1, n1 / n0 if n0 else 0.0, R0 / n0 if n0 else 0.0, R1 / n1 if n1 else 0.0)
        return out
    SUMM = [(g, keys, summary(keys)) for g, keys in GRP]
    summ = SUMM[0][2]; sess = SUMM[1][2]; nSess = len(SUMM[1][1])
    bestA = max(((k, v) for k, v in summ.items() if k[0] == 'atr'), key=lambda kv: kv[1][0] - kv[1][1])
    sg = summ.get(('sigma', 2.5))
    mid = [sess[('atr', X)] for X in (2.0, 2.5, 3.0, 3.5, 4.0, 4.5)]; s5 = sess[('atr', 5.0)]
    r = note_rows(ws, 2, [
        'THE SETTING (indicator, strategy and EA; off by default)',
        'ACCURACY FILTERS: "Skip Signals Too Far From VWAP: Max Distance" - tick it, enter the distance and choose x ATR or Price. A signal whose candle closes farther than this from the VWAP '
        '(above or below) is skipped. Like the other filters, the skipped setup still runs invisibly, so the other signals stay exactly the same, and the table row SKIPPED BY FILTERS shows what the '
        'skipped trades would have done. Default value 5 x ATR (the cut-off that did the least harm below). MT5 EA: InpMaxVwOn, InpMaxVwDist (5.0), InpMaxVwUnit (x ATR / Price). '
        'Distance = |signal close - VWAP| on the signal candle.',
        'The test below measured every trade of every setting on that same candle, in ATR, in price and in VWAP standard deviations (the existing "Don\'t Chase" filter). A cut-off was then chosen '
        'on 2021-23 only and read on 2024-26, like the main settings search; the engine with the filter switched on gave exactly the same trades as this test.'])
    r = note_rows(ws, r, [
        'WHAT IT SHOWS',
        '1. Far from the VWAP is not where the losers are on gold. The tables in part 2 show R per trade by distance: for the profitable settings the far trades (3-5 ATR and more) were often as good '
        'as or better than the near ones; for the default settings every distance loses about the same.',
        '2. No ATR cut-off helped most settings. Over the %d profitable settings, the best ATR cut-off (%.1f ATR) gave a better R per trade in both 2021-23 and 2024-26 on %d settings and a worse one in both '
        'on %d, and the total profit fell from %+.0f R to %+.0f R because it removes %.0f%% of the trades.' % (len(prof), bestA[0][1], bestA[1][0], bestA[1][1], bestA[1][2], bestA[1][3], 100 * (1 - bestA[1][4])),
        ('2b. With the Session anchor (the default VWAP; %d of these settings) every cut-off from 2 to 4.5 ATR lowered the R per trade (%+.3f without, %+.3f to %+.3f with) '
         'and was worse in both halves on up to %d settings: the far trades were the better ones. 5 ATR did the least harm (better on %d, worse on %d, %.0f%% of the trades kept, '
         '%+.3f R per trade), so it is the default value. With the Month anchor most signals are more than 5 ATR from the VWAP, so any ATR cut-off removes most trades.' % (
            nSess, mid[0][5], min(m[6] for m in mid), max(m[6] for m in mid), max(m[1] for m in mid), s5[0], s5[1], 100 * s5[4], s5[6])),
        '3. A price cut-off is the weakest: gold\'s price and its candle ranges roughly doubled from 2021 to 2026, so a fixed $ distance means different things in different years. '
        'Cut-offs chosen in price on 2021-23 usually did worse in 2024-26 (part 3).',
        ('4. The existing Don\'t Chase filter (in sigma) at 2.5 was never worse in both halves (better on %d of %d) but removed %.0f%% of the trades for about the same R per trade, so it adds little.' % (
            sg[0], len(prof), 100 * (1 - sg[4]))) if sg else '4. -',
        '5. With the default settings (which lose after spread) a cut-off makes the total loss smaller only because fewer trades are taken: on M1 and M5 the loss per trade got worse '
        'at every ATR cut-off, on M15 and H1 it moved by at most 0.06 R either way.',
        '6. Verdict: keep the max-distance filter off on gold. If your own chart shows far signals losing on another market or timeframe, test it there with the setting and the SKIPPED BY FILTERS row.'])
    # PART 1: cut-off summary
    r = section(ws, r, 'PART 1 - every cut-off over the %d settings that are profitable in both halves (and split by VWAP anchor): how many it helps or hurts in both 2021-23 and 2024-26' % len(prof), 9)
    hdr(ws, r, ['Max distance', 'Unit - settings', 'Better R per trade in both halves (settings)', 'Worse in both halves (settings)', 'Total R without the filter',
                'Total R with the filter', 'Trades kept', 'R per trade without', 'R per trade with'], height=54); r += 1
    for g, keys, sm in SUMM:
        for (unit, X), (b, w, R0, R1, kept, e0, e1) in sm.items():
            if g != GRP[0][0] and unit != 'atr': continue
            for j, v in enumerate([X, '%s - %s (%d)' % (UNIT[unit], g, len(keys)), b, w, round(R0, 1), round(R1, 1), round(kept, 4), round(e0, 4), round(e1, 4)], 1):
                put(ws, r, j, v, '0.0' if j in (1, 5, 6) else '0.0%' if j == 7 else '0.000' if j in (8, 9) else None, False, BLUE if j != 2 else None,
                    GOOD if b > w + 2 else (BAD if w > b + 2 else None))
            r += 1
        r += 1
    # PART 2: R per trade by distance (ATR)
    r += 1; r = section(ws, r, 'PART 2 - R per trade after spread by distance of the signal close from the VWAP (in ATR), every setting', 14)
    edges = [(lo, hi) for lo, hi, _ in D[order[0]]['buckets']]
    hdr(ws, r, ['Timeframe', 'Setting', 'Trades'] + ['%.1f-%s ATR' % (lo, ('%.1f' % hi) if hi < 50 else 'up') for lo, hi in edges], height=42); r += 1
    for k in order:
        d = D[k]; vals = [k[0], '%s (%s)' % (k[1], ', '.join(T[k]['scen'][:2])), d['n']]
        cells = [(s['E'], s['n']) for _, _, s in d['buckets']]
        for j, v in enumerate(vals, 1): put(ws, r, j, v, '#,##0' if j == 3 else None, False, BLUE if j == 3 else None, wrap=(j == 2))
        for j, (e, n) in enumerate(cells, 4):
            put(ws, r, j, round(e, 3) if n >= 10 else ('-' if not n else '(%d tr)' % n), '0.000' if n >= 10 else None, False, BLUE if n >= 10 else None,
                GOOD if n >= 10 and e > 0.1 else (BAD if n >= 10 and e < -0.1 else None))
        ws.row_dimensions[r].height = 24; r += 1
    # PART 3: cut-off chosen on 2021-23, read on 2024-26
    import vwfilter as VF
    r += 1; r = section(ws, r, 'PART 3 - cut-off chosen on 2021-23 only (most R, at least 60% of trades kept), then read on 2024-26', 14)
    hdr(ws, r, ['Timeframe', 'Setting', 'Unit', 'Chosen max distance', 'Trades kept', '2021-23 R without', '2021-23 R with', '2024-26 R without', '2024-26 R with',
                '2024-26 R per trade without', '2024-26 R per trade with', 'Longest losing streak without', 'Longest losing streak with', '2024-26 better?'], height=54); r += 1
    for k in order:
        d = D[k]; b = d['res']['none']
        for unit in ('atr', 'sigma', 'price'):
            pk = VF.pick(d['res'], unit)
            if not pk: continue
            (u, X), v = pk
            ok = v['B']['E'] > b['B']['E']
            vals = [k[0], k[1], UNIT[unit], X, round(v['all']['n'] / max(1, b['all']['n']), 4), round(b['A']['R'], 1), round(v['A']['R'], 1), round(b['B']['R'], 1), round(v['B']['R'], 1),
                    round(b['B']['E'], 4), round(v['B']['E'], 4), b['all']['maxL'], v['all']['maxL'], 'yes' if ok else 'no']
            for j, val in enumerate(vals, 1):
                put(ws, r, j, val, '0.0%' if j == 5 else '0.0' if 6 <= j <= 9 else '0.000' if j in (10, 11) else None, False, BLUE if 4 <= j <= 13 else None,
                    GOOD if j == 14 and ok else (BAD if j == 14 else None))
            r += 1
        r += 1
    widths(ws, [9, 44, 16, 11, 11, 11, 11, 11, 11, 11, 11, 11, 11, 10])
    ws.freeze_panes = 'C1'
