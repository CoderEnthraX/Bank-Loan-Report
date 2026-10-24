"""Final workbook: best settings, summaries, year x month sheets (default / best / low-streak) per timeframe, TP1-TP4 detail,
settings sweep, combinations, exits, recovery rules, streak math. Counts are backtest results (blue); rates and totals are formulas."""
import sys, os, pickle, collections, math, statistics
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter as CL
from openpyxl.comments import Comment
S = '/tmp/claude-0/-home-user-Bank-Loan-Report/21ed2550-acec-5cf1-bf01-a24cbf22fe18/scratchpad'
sys.path.insert(0, S + '/gold')
import geng, greedy, labels, anaA, recovery as RC
F = 'Arial'
def _days():
    d = pickle.load(open(S + '/gold/data/H1.pkl', 'rb')); ds = set(t.date() for t in d['srv']); c = collections.Counter(x.year for x in ds)
    return len(ds), dict(c)
DAYS_N, DAYS_Y = _days()
HDR = PatternFill('solid', fgColor='1F3864'); YR = PatternFill('solid', fgColor='D9E1F2'); TOT = PatternFill('solid', fgColor='FFE699')
SUB = PatternFill('solid', fgColor='E2EFDA'); BAD = PatternFill('solid', fgColor='F8CBAD'); GOOD = PatternFill('solid', fgColor='C6EFCE')
thin = Side(style='thin', color='BFBFBF'); BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
MON = ['Jan', 'Feb', 'Mar', 'Apr', 'May', 'Jun', 'Jul', 'Aug', 'Sep', 'Oct', 'Nov', 'Dec']
TFS = ('M1', 'M5', 'M15', 'H1')
PROF = (('default', 'default'), ('profit', 'best'), ('streak', 'low-streak'))
PNAME = {'default': 'Default settings', 'profit': 'Best settings (profit)', 'streak': 'Best settings with short losing streaks'}
BLUE = '0000FF'; GREEN = '008000'
def font(bold=False, color=None, size=9, italic=False): return Font(name=F, bold=bold, color=color, size=size, italic=italic)
def hdr(ws, r, cols, c0=1, height=42):
    for j, c in enumerate(cols, c0):
        cell = ws.cell(r, j, c); cell.font = Font(name=F, bold=True, color='FFFFFF', size=9); cell.fill = HDR
        cell.alignment = Alignment(wrap_text=True, horizontal='center', vertical='center'); cell.border = BOX
    ws.row_dimensions[r].height = height
def put(ws, r, c, v, fmt=None, bold=False, color=None, fill=None, italic=False, wrap=False, align=None):
    cell = ws.cell(r, c, v); cell.font = font(bold, color, italic=italic); cell.border = BOX
    if fmt: cell.number_format = fmt
    if fill: cell.fill = fill
    if wrap or align: cell.alignment = Alignment(wrap_text=wrap, horizontal=align, vertical='top' if wrap else None)
    return cell
def title(ws, text, sub=None):
    ws['A1'] = text; ws['A1'].font = Font(name=F, bold=True, size=12)
    if sub: ws['A2'] = sub; ws['A2'].font = font(italic=True)
def widths(ws, w):
    for j, x in enumerate(w, 1): ws.column_dimensions[CL(j)].width = x
def streaks(seq):
    mw, ml, _ = geng.streaks(seq); return mw, ml
def runs(bools):
    bt = bf = ct = cf = 0
    for b in bools:
        if b: ct += 1; cf = 0; bt = max(bt, ct)
        else: cf += 1; ct = 0; bf = max(bf, cf)
    return bt, bf
# ---------------------------------------------------------------- year x month sheet
COLS = ['Year', 'Month', 'Trades', 'Wins (closed at final target)', 'Profitable trades', 'TP1 hits', 'TP2 hits', 'TP3 hits', 'TP4 hits', 'Breakeven exits',
        'Total R', 'Spread cost (R)', 'Total R after spread', 'Longest win streak', 'Longest losing streak',
        'Win %', 'Profitable %', 'TP1 %', 'TP2 %', 'TP3 %', 'TP4 %', 'R per trade', 'R per trade after spread']
def counts(tr):
    res = [x[1] for x in tr]; mw, ml = streaks(res)
    return dict(n=len(tr), w=sum(1 for x in tr if x[2]), p=sum(1 for x in tr if x[1] > 0), h=[sum(1 for x in tr if x[3][k]) for k in range(4)],
                be=sum(1 for x in tr if x[6]), R=sum(res), cost=sum(x[4] / x[5] for x in tr), mw=mw, ml=ml)
def ym_row(ws, row, y, m, c, fill=None, bold=False, sumrange=None):
    put(ws, row, 1, y, bold=bold, fill=fill); put(ws, row, 2, m, bold=bold, fill=fill)
    if sumrange is None:
        for j, v in enumerate([c['n'], c['w'], c['p'], *c['h'], c['be'], round(c['R'], 4), round(c['cost'], 4)], 3):
            put(ws, row, j, v, '0.00' if j in (11, 12) else None, bold, BLUE, fill)
    else:
        for j in range(3, 13): put(ws, row, j, sumrange(CL(j)), '0.00' if j in (11, 12) else None, bold, None, fill)
    put(ws, row, 13, '=K%d-L%d' % (row, row), '0.00', bold, None, fill)
    put(ws, row, 14, c['mw'], None, bold, BLUE, fill); put(ws, row, 15, c['ml'], None, bold, BLUE, fill)
    for j, src in zip(range(16, 22), 'DEFGHI'):
        put(ws, row, j, '=IF($C%d=0,"",%s%d/$C%d)' % (row, src, row, row), '0.0%', bold, None, fill)
    put(ws, row, 22, '=IF($C%d=0,"",K%d/$C%d)' % (row, row, row), '0.000', bold, None, fill)
    put(ws, row, 23, '=IF($C%d=0,"",M%d/$C%d)' % (row, row, row), '0.000', bold, None, fill)
def ym_sheet(wb, name, tr, head, cfg):
    ws = wb.create_sheet(name); title(ws, head)
    ws['A2'] = ('Settings: ' + labels.short(cfg)); ws['A2'].font = font(italic=True); ws.merge_cells('A2:W2')
    ws['A2'].alignment = Alignment(wrap_text=True, vertical='top'); ws.row_dimensions[2].height = 42
    ws['A3'] = ('Blue = backtest counts / R / streaks; black = formulas. Month = IST date of the fill. Monthly streaks are counted inside the month; '
                'year and 2021-2026 streaks use the full trade sequence (a breakeven exit ends a streak).')
    ws['A3'].font = font(italic=True, size=8); ws.merge_cells('A3:W3')
    hdr(ws, 5, COLS); by = collections.defaultdict(lambda: collections.defaultdict(list))
    for x in tr: by[x[0].year][x[0].month].append(x)
    r = 6; yrows = {}
    empty = dict(n=0, w=0, p=0, h=[0, 0, 0, 0], be=0, R=0.0, cost=0.0, mw=0, ml=0)
    for y in sorted(by):
        first = r
        for m in range(1, 13):
            if y == 2026 and m > 10: break
            ym_row(ws, r, y, MON[m - 1], counts(by[y][m]) if by[y].get(m) else empty); r += 1
        c = counts([x for m in sorted(by[y]) for x in by[y][m]])
        ym_row(ws, r, '%d total' % y, 'All', c, YR, True, lambda L, a=first, b=r - 1: '=SUM(%s%d:%s%d)' % (L, a, L, b))
        yrows[y] = r; r += 1
    ca = counts(tr)
    ym_row(ws, r, '2021-2026', 'All', ca, TOT, True, lambda L: '=' + '+'.join('%s%d' % (L, k) for k in yrows.values()))
    widths(ws, [13, 7, 7, 9, 9, 7, 7, 7, 7, 9, 8, 8, 9, 8, 8, 7, 9, 7, 7, 7, 7, 8, 9]); ws.freeze_panes = 'C6'
    return yrows, r
# ---------------------------------------------------------------- summary
def summary(wb, refs):
    ws = wb.create_sheet('Summary', 2)
    title(ws, 'Accuracy per timeframe, settings and year (FxPro MT5 GOLD, 2021-2026)', 'Every value links to the year-total rows of the year x month sheets (green).')
    cols = ['Timeframe', 'Settings', 'Year', 'Trades', 'Trades per day', 'Win %', 'Profitable %', 'TP1 %', 'TP2 %', 'TP3 %', 'TP4 %', 'Total R after spread',
            'R per trade', 'R per trade after spread', 'Longest win streak', 'Longest losing streak']
    hdr(ws, 4, cols); r = 5
    for tf, prof, sheet, (yrows, tot) in refs:
        for y, row in list(yrows.items()) + [('2021-2026', tot)]:
            q = "'%s'!" % sheet; last = y == '2021-2026'; fl = TOT if last else None
            put(ws, r, 1, tf, bold=last, fill=fl); put(ws, r, 2, PNAME[prof], bold=last, fill=fl); put(ws, r, 3, str(y), bold=last, fill=fl)
            put(ws, r, 4, '=%sC%d' % (q, row), None, last, GREEN, fl)
            put(ws, r, 5, '=D%d/%d' % (r, DAYS_N if last else DAYS_Y[y]), '0.00', last, None, fl)
            for j, src in zip(range(6, 12), 'PQRSTU'): put(ws, r, j, '=%s%s%d' % (q, src, row), '0.0%', last, GREEN, fl)
            put(ws, r, 12, '=%sM%d' % (q, row), '0.0', last, GREEN, fl); put(ws, r, 13, '=%sV%d' % (q, row), '0.000', last, GREEN, fl)
            put(ws, r, 14, '=%sW%d' % (q, row), '0.000', last, GREEN, fl)
            put(ws, r, 15, '=%sN%d' % (q, row), None, last, GREEN, fl); put(ws, r, 16, '=%sO%d' % (q, row), None, last, GREEN, fl)
            r += 1
        r += 1
    ws.cell(3, 1, 'Trades per day = trades / trading days in the data (2026 = 2 Jan - 2 Oct). Year rows: ' + ', '.join('%d: %d days' % (y, d) for y, d in sorted(DAYS_Y.items())) + '.').font = font(italic=True, size=8)
    widths(ws, [9, 30, 10, 8, 8, 7, 9, 7, 7, 7, 7, 10, 8, 10, 8, 9]); ws.freeze_panes = 'D5'
# ---------------------------------------------------------------- TP1-TP4 detail
def tp_sheet(wb, FIN):
    ws = wb.create_sheet('TP1-TP4')
    title(ws, 'TP1, TP2, TP3 and TP4 individually: accuracy and win / losing streaks',
          'Part 1: with the settings as they are, how often each target was reached before the stop, and the longest run of trades that did / did not reach it. '
          'Part 2: each target used as the real final exit (TP1 = 0.6R, TP2 = 1R, TP3 = 1.5R, TP4 = 2R), same signals, no stop moves, no partials.')
    ws.merge_cells('A2:P2'); ws['A2'].alignment = Alignment(wrap_text=True); ws.row_dimensions[2].height = 30
    r = 4; put(ws, r, 1, 'PART 1 - targets reached (with the final exit of each setting)', bold=True, fill=SUB); r += 1
    cols = ['Timeframe', 'Settings', 'Year', 'Trades'] + sum([['TP%d hits' % k, 'TP%d %%' % k, 'TP%d longest run reached' % k, 'TP%d longest run not reached' % k] for k in (1, 2, 3, 4)], [])
    hdr(ws, r, cols, height=54); r += 1
    for tf in TFS:
        if tf not in FIN: continue
        for prof, _ in PROF:
            tr = FIN[tf]['trades'][prof]; ex = dict(greedy.defaults()); ex.update(FIN[tf]['cfg'][prof]); xk = ex['exitAt']
            by = collections.defaultdict(list)
            for x in tr: by[x[0].year].append(x)
            for y in sorted(by) + ['2021-2026']:
                q = tr if y == '2021-2026' else by[y]; last = y == '2021-2026'; fl = TOT if last else None
                put(ws, r, 1, tf, bold=last, fill=fl); put(ws, r, 2, PNAME[prof], bold=last, fill=fl); put(ws, r, 3, str(y), bold=last, fill=fl)
                put(ws, r, 4, len(q), None, last, BLUE, fl)
                for k in range(4):
                    c0 = 5 + 4 * k
                    if k + 1 > xk:
                        for j in range(4): put(ws, r, c0 + j, 'closes at TP%d' % xk if j == 0 else '-', None, last, None, fl)
                        continue
                    hb = [x[3][k] for x in q]; w, l = runs(hb)
                    put(ws, r, c0, sum(hb), None, last, BLUE, fl)
                    put(ws, r, c0 + 1, '=IF($D%d=0,"",%s%d/$D%d)' % (r, CL(c0), r, r), '0.0%', last, None, fl)
                    put(ws, r, c0 + 2, w, None, last, BLUE, fl); put(ws, r, c0 + 3, l, None, last, BLUE, fl)
                r += 1
            r += 1
    r += 1; put(ws, r, 1, 'PART 2 - each target as the real final exit (2021-2026, after spread)', bold=True, fill=SUB); r += 1
    cols = ['Timeframe', 'Signals from', 'Final exit', 'Trades', 'Wins', 'Win %', 'Total R after spread', 'R per trade after spread', '2021-23 R per trade after spread',
            '2024-26 R per trade after spread', 'Years positive (of 6)', 'Longest win streak', 'Longest losing streak', 'Worst single-year losing streak']
    hdr(ws, r, cols, height=54); r += 1
    for tf in TFS:
        if tf not in FIN: continue
        for prof, _ in PROF:
            for k in (1, 2, 3, 4):
                cfg, tr = FIN[tf]['exits'][(prof, k)]
                res = [x[1] for x in tr]; mw, ml = streaks(res); net = [x[1] - x[4] / x[5] for x in tr]
                A = [v for x, v in zip(tr, net) if x[0].year <= 2023]; B = [v for x, v in zip(tr, net) if x[0].year >= 2024]
                yrs = collections.defaultdict(list)
                for x, v in zip(tr, net): yrs[x[0].year].append((x[1], v))
                dd = dict(greedy.defaults()); dd.update(cfg)
                put(ws, r, 1, tf); put(ws, r, 2, PNAME[prof]); put(ws, r, 3, 'TP%d (%sR)' % (k, labels.fmt(float(dd['tp%dR' % k]))))
                put(ws, r, 4, len(tr), None, False, BLUE); put(ws, r, 5, sum(1 for x in tr if x[2]), None, False, BLUE)
                put(ws, r, 6, '=IF(D%d=0,"",E%d/D%d)' % (r, r, r), '0.0%')
                put(ws, r, 7, round(sum(net), 3), '0.0', False, BLUE); put(ws, r, 8, '=IF(D%d=0,"",G%d/D%d)' % (r, r, r), '0.000')
                put(ws, r, 9, round(sum(A) / len(A), 4) if A else 0, '0.000', False, BLUE); put(ws, r, 10, round(sum(B) / len(B), 4) if B else 0, '0.000', False, BLUE)
                put(ws, r, 11, sum(1 for y in yrs if sum(v for _, v in yrs[y]) > 0), None, False, BLUE)
                put(ws, r, 12, mw, None, False, BLUE); put(ws, r, 13, ml, None, False, BLUE)
                put(ws, r, 14, max(streaks([a for a, _ in yrs[y]])[1] for y in yrs), None, False, BLUE)
                r += 1
            r += 1
    widths(ws, [9, 30, 10, 7] + [7, 7, 9, 9] * 4); ws.freeze_panes = 'D6'
# ---------------------------------------------------------------- settings sweep
def sweep_sheet(wb):
    ws = wb.create_sheet('Settings sweep')
    title(ws, 'Every indicator setting changed one at a time (all others default), final exit TP4, after spread',
          'Sorted by 2021-23 result. 2021-23 = where settings were judged; 2024-26 = the check. Blue = backtest results; R per trade = formula.')
    r = 4
    cols = ['Timeframe', 'Setting changed', 'Trades', '2021-23 trades', '2021-23 total R after spread', '2021-23 R per trade', '2024-26 trades', '2024-26 total R after spread',
            '2024-26 R per trade', 'Years positive', 'Win %', 'TP1 %', 'TP2 %', 'TP3 %', 'TP4 %', 'Longest win streak', 'Longest losing streak']
    hdr(ws, r, cols, height=54); r += 1
    for tf in TFS:
        if not (os.path.exists(S + '/gold/sweep_%s.pkl' % tf) or os.path.exists(S + '/gold/slim_sweep_%s.pkl' % tf)): continue
        C, M = anaA.sweep_metrics(tf); rows = [(n, M[n]) for n, _ in C if n in M]
        rows = [x for x in rows if x[1]]; rows.sort(key=lambda x: (x[0] != 'default', -x[1]['EcA']))
        for n, m in rows:
            fl = TOT if n == 'default' else None
            vals = [tf, n, m['n'], m['nA'], round(m['EcA'] * m['nA'], 3), '=IF(D{0}=0,"",E{0}/D{0})', m['nB'], round(m['EcB'] * m['nB'], 3), '=IF(G{0}=0,"",H{0}/G{0})',
                    m['ypos'], round(m['win'], 4)] + [round(t['rate'], 4) for t in m['tps']] + [m['maxW'], m['maxL']]
            for j, v in enumerate(vals, 1):
                isf = isinstance(v, str) and v.startswith('=')
                put(ws, r, j, v.format(r) if isf else v, '0.000' if j in (6, 9) else '0.0' if j in (5, 8) else '0.0%' if 11 <= j <= 15 else None, n == 'default',
                    None if (isf or j <= 2) else BLUE, fl)
            r += 1
        r += 1
    widths(ws, [9, 46, 7, 8, 10, 9, 8, 10, 9, 8, 7, 7, 7, 7, 7, 8, 8]); ws.freeze_panes = 'C5'
# ---------------------------------------------------------------- combinations + search paths
def combo_sheet(wb):
    ws = wb.create_sheet('Combinations')
    title(ws, 'Settings in combination: the best pairs / triples / quadruples, and the step-by-step search over every setting',
          'Part 1: combinations (size 1-4) of the settings that helped alone. Part 2: the search - each step adds the one change (of ~170) that most improves 2021-23; 2024-26 is only reported.')
    ws.merge_cells('A2:N2'); ws['A2'].alignment = Alignment(wrap_text=True); ws.row_dimensions[2].height = 28
    r = 4; put(ws, r, 1, 'PART 1 - combinations (final exit TP4), top 25 per timeframe by 2021-23', bold=True, fill=SUB); r += 1
    cols = ['Timeframe', 'Combination', 'Trades', '2021-23 R per trade after spread', '2024-26 R per trade after spread', 'Years positive', 'Win %', 'Longest win streak', 'Longest losing streak']
    hdr(ws, r, cols, height=54); r += 1
    for tf in TFS:
        fn = S + '/gold/stageB_%s.pkl' % tf; fs = S + '/gold/slim_stageB_%s.pkl' % tf
        if os.path.exists(fn): D = pickle.load(open(fn, 'rb')); rows = [(n, anaA.metrics(t)) for n, t in D['res'].items() if t]
        elif os.path.exists(fs): rows = list(pickle.load(open(fs, 'rb'))['metrics'].items())
        else: continue
        d0 = anaA.sweep_metrics(tf)[1]['default']; rows = [x for x in rows if x[1] and x[1]['nA'] >= 0.15 * d0['nA']]
        rows.sort(key=lambda x: -x[1]['EcA'])
        pos = [x for x in rows if x[1]['EcA'] > 0]
        for n, m in rows[:25]:
            for j, v in enumerate([tf, n, m['n'], round(m['EcA'], 4), round(m['EcB'], 4), m['ypos'], round(m['win'], 4), m['maxW'], m['maxL']], 1):
                put(ws, r, j, v, '0.000' if j in (4, 5) else '0.0%' if j == 7 else None, False, BLUE if j > 2 else None,
                    (GOOD if v > 0 else BAD) if j in (4, 5) else None, wrap=(j == 2))
            r += 1
        put(ws, r, 2, '%s: %d of %d combinations positive in 2021-23 after spread; %d of those also positive in 2024-26.' % (tf, len(pos), len(rows), sum(1 for x in pos if x[1]['EcB'] > 0)), italic=True)
        r += 2
    r += 1; put(ws, r, 1, 'PART 2 - step-by-step search over every setting (chosen on 2021-23 only)', bold=True, fill=SUB); r += 1
    cols = ['Timeframe', 'Search', 'Step', 'Change added', '2021-23 trades', '2021-23 R per trade after spread', '2024-26 trades', '2024-26 R per trade after spread',
            'Years positive', '2021-23 win %', '2021-23 longest W', '2021-23 longest L', '2024-26 longest W', '2024-26 longest L']
    hdr(ws, r, cols, height=54); r += 1
    SN = {'profit:default': 'Most profit, from default', 'profit:bestB': 'Most profit, from best combination', 'streak:default': 'Losing streak <= 7, from default',
          'streak:profit': 'Losing streak <= 7, from the most-profit result'}
    for tf in TFS:
        fn = S + '/gold/greedy_%s.pkl' % tf
        if not os.path.exists(fn): continue
        P = pickle.load(open(fn, 'rb'))
        for sid, path in P.items():
            for k, st in enumerate(path):
                m = st['m']; mv = 'start: ' + (' + '.join(st['names']) or 'default') if k == 0 else st['move']
                vals = [tf, SN.get(sid, sid), k, mv, m['A']['n'], round(m['A']['E'], 4), m['B']['n'], round(m['B']['E'], 4), m['ypos'], round(m['A']['win'], 4),
                        m['A']['maxW'], m['A']['maxL'], m['B']['maxW'], m['B']['maxL']]
                for j, v in enumerate(vals, 1):
                    put(ws, r, j, v, '0.000' if j in (6, 8) else '0.0%' if j == 10 else None, False, BLUE if j > 4 else None,
                        (GOOD if v > 0 else BAD) if j in (6, 8) else None, wrap=(j == 4))
                r += 1
            r += 1
    widths(ws, [9, 34, 6, 46, 8, 10, 8, 10, 8, 8, 8, 8, 8, 8]); ws.freeze_panes = 'C4'
# ---------------------------------------------------------------- exits (stage C)
EXN = ['Exit at TP1 (0.6R)', 'Exit at TP2 (1R)', 'Exit at TP3 (1.5R)', 'Exit at TP4 (2R)', 'TP4 + breakeven at TP1', 'TP2 + breakeven at TP1', 'TP4 + breakeven at TP2', '50% at TP1 + breakeven, rest TP4']
def exm(tr):
    n = len(tr); res = [x[1] for x in tr]; mw, ml = streaks(res)
    A = [x for x in tr if x[0].year <= 2023]; B = [x for x in tr if x[0].year >= 2024]
    ec = lambda q: sum(x[1] - x[4] / x[5] for x in q) / len(q) if q else 0.0
    yrs = collections.defaultdict(list)
    for x in tr: yrs[x[0].year].append(x)
    return dict(n=n, win=sum(1 for v in res if v > 0) / n, EA=ec(A), EB=ec(B), Ec=ec(tr), Rc=sum(x[1] - x[4] / x[5] for x in tr),
                ypos=sum(1 for y in yrs if sum(x[1] - x[4] / x[5] for x in yrs[y]) > 0), mw=mw, ml=ml)
def exits_sheet(wb):
    ws = wb.create_sheet('Exits')
    title(ws, 'Exit choices for the candidate settings: each TP as the exit, breakeven moves, 50% partial at TP1',
          'Breakeven = stop moved to the entry when that target is hit (a breakeven exit is neither a win nor a loss and ends a streak). After spread.')
    r = 4
    cols = ['Timeframe', 'Candidate settings', 'Exit', 'Trades', 'Trades per day', 'Win %', 'R per trade after spread', '2021-23', '2024-26', 'Years positive',
            'Longest win streak', 'Longest losing streak', 'Total R after spread']
    hdr(ws, r, cols, height=42); r += 1
    for tf in TFS:
        fn = S + '/gold/stageC_%s.pkl' % tf; fs = S + '/gold/slim_stageC_%s.pkl' % tf
        if os.path.exists(fn): D = pickle.load(open(fn, 'rb')); get = lambda k: exm(D['res'][k])
        elif os.path.exists(fs): D = pickle.load(open(fs, 'rb')); get = lambda k: D['metrics'][k]
        else: continue
        for cn, _ in D['cands']:
            for e in EXN:
                q = get((cn, e))
                vals = [tf, cn, e, q['n'], '=D%d/%d' % (r, DAYS_N), round(q['win'], 4), round(q['Ec'], 4), round(q['EA'], 4), round(q['EB'], 4), q['ypos'], q['mw'], q['ml'], round(q['Rc'], 2)]
                for j, v in enumerate(vals, 1):
                    put(ws, r, j, v, '0.00' if j == 5 else '0.0%' if j == 6 else '0.000' if j in (7, 8, 9) else '0.0' if j == 13 else None, False,
                        None if j in (1, 2, 3, 5) else BLUE, (GOOD if v > 0 else BAD) if j in (7, 8, 9) else None, wrap=(j == 2))
                r += 1
            r += 1
    widths(ws, [9, 44, 30, 7, 8, 7, 9, 8, 8, 8, 8, 8, 9]); ws.freeze_panes = 'D5'
# ---------------------------------------------------------------- recovery
def rkey(x): return (x['mode'], round(x['net'], 2), round(x['maxDD'], 2))
def eligible(rows):
    """survived, made money, never risked more than $500 on one trade, never fell more than 40% from a peak; a capped run identical to its uncapped run counts once"""
    seen = set(); out = []
    for x in sorted([x for x in rows if not x['ruined'] and x['net'] > 0 and x['maxDD'] > 0 and x['maxRisk'] <= 500.01 and x['maxDDpct'] <= 0.40],
                    key=lambda x: (-x['net'] / x['maxDD'], x['cap'])):
        if rkey(x) in seen: continue
        seen.add(rkey(x)); out.append(x)
    return out
RNOTE = {'Off': 'fixed $50 risk', 'Recover all losses on the next trade': 'risk = $50 + all open losses', 'Recover losses over N trades': 'N = 3',
         'Recover a % of losses per trade': '50% of open losses', 'Martingale (multiply after each loss)': 'x2 after each loss', 'Exact recovery at final TP': 'sized so the final TP repays open losses',
         'Step-up (+$ per loss in a row)': '+$25 per loss in a row', 'Fibonacci': '1,1,2,3,5... units', "D'Alembert": '+1 unit after a loss, -1 after a win',
         'Rule C: carry the losses': 'carried losses / 3', 'Rule C split': 'Rule C with profit steps', 'Labouchere': '1-2-3 line', 'Grand Martingale': 'x2 + 1 unit',
         'Paroli': 'x2 after a win, 3-win reset', "Reverse D'Alembert": '+1 unit after a win', 'Delayed recovery': 'starts after 2 losses', '% of equity': '1% of equity',
         "Oscar's Grind": '+1 unit after a win until +1 unit up'}
def recovery_sheet(wb, FIN):
    ws = wb.create_sheet('Recovery rules')
    title(ws, 'All 17 loss-recovery rules of the strategy + Off, on each timeframe and setting (2021-2026, after spread)',
          '$10,000 start, $50 base risk, sized like the strategy (a stop-out loses the computed risk incl. spread), no leverage limit. "$500 cap" = Hard Cap on, Clamp. '
          'Ranked only if the account survived, made money, never risked more than $500 (5% of the start) on one trade and never fell more than 40% from a peak.')
    ws.merge_cells('A2:N2'); r = 4
    # cross-check: every rule on every profitable pick (default settings excluded: every recovery rule blows the account there)
    agg = collections.defaultdict(list); surv = collections.Counter(); cases = 0
    for tf in TFS:
        if tf not in FIN: continue
        for prof in ('profit', 'streak'):
            if FIN[tf].get('passed', {}).get(prof) is False: continue
            rows = FIN[tf]['recovery'][prof]; cases += 1; off = [x for x in rows if x['mode'] == 'Off'][0]
            el = eligible(rows); rank = {rkey(x): k + 1 for k, x in enumerate(el)}
            for x in rows:
                if rkey(x) in rank: agg[(x['mode'], x['cap'])].append((rank[rkey(x)], (x['net'] / x['maxDD']) / (off['net'] / off['maxDD']), x['net'] / off['net']))
            for x in rows: surv[(x['mode'], x['cap'])] += 0 if x['ruined'] else 1
    put(ws, r, 1, 'SUMMARY - every rule on the %d profitable best / low-streak settings (all timeframes)' % cases, bold=True, fill=SUB)
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=9); r += 1
    hdr(ws, r, ['Recovery rule', 'Hard cap', 'Account survived (of %d)' % cases, 'Eligible (survived, profitable, max risk <= $500, drawdown <= 40%)', 'Average rank (1 = best)',
                'Profit / drawdown vs Off (average)', 'Net profit vs Off (average)', 'Verdict'], height=54); r += 1
    order = sorted(agg.items(), key=lambda kv: (-len(kv[1]), -sum(q for _, q, _ in kv[1]) / len(kv[1])))
    for (mode, cap), v in order:
        q = sum(x for _, x, _ in v) / len(v); g = sum(x for _, _, x in v) / len(v)
        verdict = ('baseline' if mode == 'Off' else 'about as good as Off, more profit' if q >= 0.85 else 'worse risk-adjusted than Off' if q >= 0.7 else 'much worse than Off')
        for j, val in enumerate([mode, '$500' if cap else 'none', surv[(mode, cap)], len(v), round(sum(x for x, _, _ in v) / len(v), 1), round(q, 3), round(g, 2), verdict], 1):
            put(ws, r, j, val, '0.00' if j in (6, 7) else None, False, BLUE if 3 <= j <= 7 else None, GOOD if mode == 'Off' else None)
        r += 1
    ws.cell(r, 1, 'Rules that were never eligible (blew the account, needed more than $500 on one trade or fell more than 40% in every case) are not listed. Detail per timeframe below.').font = font(italic=True, size=8)
    r += 2
    cols = ['Timeframe', 'Settings', 'Recovery rule', 'Rule detail', 'Hard cap', 'Net profit ($)', 'Max drawdown ($)', 'Max drawdown %', 'Largest risk on one trade ($)',
            'Account blown?', 'Profit / max drawdown', 'Profit factor', 'Worst run of losing trades ($)', 'Rank (survived, by profit / drawdown)']
    hdr(ws, r, cols, height=54); r += 1; best = []
    for tf in TFS:
        if tf not in FIN: continue
        for prof, _ in PROF:
            rows = FIN[tf]['recovery'][prof]
            alive = eligible(rows); rk0 = {rkey(x): k + 1 for k, x in enumerate(alive)}; rk = {id(x): rk0[rkey(x)] for x in rows if rkey(x) in rk0}
            if alive: best.append((tf, prof, alive[0], [x for x in rows if x['mode'] == 'Off' and x['cap'] == 0][0]))
            for x in rows:
                fl = GOOD if alive and rkey(x) == rkey(alive[0]) else (BAD if x['ruined'] else None)
                vals = [tf, PNAME[prof], x['mode'], RNOTE.get(x['mode'], ''), 'none' if x['cap'] == 0 else '$500', round(x['net'], 2), round(x['maxDD'], 2), round(x['maxDDpct'], 4),
                        round(x['maxRisk'], 2), ('YES, after %d trades' % x['ruinAt']) if x['ruined'] else 'no', '=IF(G{0}=0,"",F{0}/G{0})',
                        round(x['pf'], 3) if x['pf'] != float('inf') else 'n/a', round(x['worstRun'], 2), rk.get(id(x), '-')]
                for j, v in enumerate(vals, 1):
                    isf = isinstance(v, str) and v.startswith('=')
                    put(ws, r, j, v.format(r) if isf else v, '#,##0;(#,##0);-' if j in (6, 7, 9, 13) else '0.0%' if j == 8 else '0.00' if j in (11, 12) else None,
                        False, BLUE if (5 < j < 14 and not isf and j != 10) else None, fl)
                r += 1
            r += 1
    widths(ws, [9, 30, 34, 30, 8, 11, 11, 9, 10, 18, 9, 8, 11, 10]); ws.freeze_panes = 'D5'
    return best
# ---------------------------------------------------------------- Rule C and Rule C split with the loss split over 1-5 trades
def rulec_sheet(wb):
    fn = S + '/gold/rulec.pkl'
    if not os.path.exists(fn): return
    D = pickle.load(open(fn, 'rb'))['rows']
    ws = wb.create_sheet('Rule C split test')
    title(ws, 'Rule C and Rule C split with the loss split over 1, 2, 3, 4 or 5 trades (2021-2026, after spread)',
          'Rule C: carried = how far closed equity is below its last high; next risk = max(base, carried / split). Rule C split: the same, and the base steps up with total profit '
          '(200 -> 200/3, 500 -> 500/5, 1000 -> 1000/8, then +1 part per +500). $10,000 start, $50 base, hard cap $500 with Clamp unless shown. Partial closes update the carried losses per part.')
    ws.merge_cells('A2:Q2'); ws['A2'].alignment = Alignment(wrap_text=True, vertical='top'); ws.row_dimensions[2].height = 40
    notes = ['WHAT IT SHOWS',
             '1. With the best settings (profitable signals) and the hard cap on, Rule C split at split 2 or 3 made 5-9x the profit of fixed risk (Off), and per $ staked it lost about the same as Off.',
             '2. Without the hard cap, Rule C split at split 1-3 blew the $10,000 account on M1, M5 and M15. With the default settings every Rule C variant blew the account.',
             '3. Split 3 was the better choice than 2 on most cases: similar profit with a smaller drawdown (M15 best: 33% vs 46%). Drawdowns are still far larger than Off (16-70% vs 3-17%).',
             '4. The cap action matters: "Start again" (forget the carried losses when the risk would pass the cap) kept M5 alive where Clamp blew it, and gave M15 its best profit for the drawdown.',
             '5. Stress test (part 3): add 0.05-0.10 R of extra cost per trade (live slippage, weaker edge). H1 stays clearly profitable; M15 holds to about 0.05 R; M1 and M5 blow up - their edge is too thin.']
    r = 4
    for k, t in enumerate(notes): ws.cell(r + k, 1, t).font = font(bold=(k == 0), size=9)
    r += len(notes) + 1
    put(ws, r, 1, 'PART 1 - split 1 to 5, without and with the $500 hard cap (Clamp)', bold=True, fill=SUB); ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=10); r += 1
    cols = ['Timeframe', 'Settings', 'Rule', 'Split', 'Hard cap', 'Net profit ($)', 'Max drawdown ($)', 'Max drawdown %', 'Average stake ($)', 'Largest stake ($)', 'Account blown?',
            'Profit / max drawdown', 'vs Off (1 = same)', 'Profit per $ of average stake', 'Max drawdown per $ of average stake', '2021-23 alone: net ($)', '2024-26 alone: net ($)']
    hdr(ws, r, cols, height=54); r += 1
    for tf in TFS:
        for prof, _ in PROF:
            if (tf, prof) not in D: continue
            offrow = None
            for x in D[(tf, prof)]:
                a, A, B = x['all'], x['A'], x['B']; off = x['mode'] == 'Off'
                rule = 'Off (fixed $50)' if off else ('Rule C' if x['mode'].startswith('Rule C:') else 'Rule C split')
                if off: offrow = r
                fl = TOT if off else (BAD if a['ruined'] else None)
                vals = [tf, PNAME[prof], rule, '-' if off else x['split'], '-' if off else ('$500' if x['cap'] else 'none'), round(a['net'], 2), round(a['maxDD'], 2), round(a['maxDDpct'], 4),
                        round(a['avgRisk'], 2), round(a['maxRisk'], 2), ('YES, after %d trades' % a['ruinAt']) if a['ruined'] else 'no',
                        '=IF(G{0}=0,"",F{0}/G{0})', '=IF(OR(L{1}="",L{1}<=0,L{0}=""),"",L{0}/L{1})' if offrow and not off else ('1' if off else ''),
                        '=IF(I{0}=0,"",F{0}/I{0})', '=IF(I{0}=0,"",G{0}/I{0})',
                        'BLOWN' if A['ruined'] else round(A['net'], 2), 'BLOWN' if B['ruined'] else round(B['net'], 2)]
                for j, v in enumerate(vals, 1):
                    isf = isinstance(v, str) and v.startswith('=')
                    put(ws, r, j, v.format(r, offrow) if isf else (1 if v == '1' else v),
                        '#,##0;(#,##0);-' if j in (6, 7, 9, 10, 16, 17) else '0.0%' if j == 8 else '0.00' if j in (12, 13) else '0.0' if j in (14, 15) else None,
                        off, BLUE if (6 <= j <= 10 or j in (16, 17)) and not isf else None, fl)
                r += 1
            r += 1
    fc = S + '/gold/rulec_caps.pkl'
    if os.path.exists(fc):
        C = pickle.load(open(fc, 'rb'))
        r += 1; put(ws, r, 1, 'PART 2 - split 2 and 3: hard cap size and what happens when the risk would pass it (Clamp = trade at the cap, Start again = forget the carried losses)',
                     bold=True, fill=SUB); ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=12); r += 1
        hdr(ws, r, ['Timeframe', 'Settings', 'Rule', 'Split', 'Hard cap', 'If the risk would pass the cap', 'Net profit ($)', 'Max drawdown ($)', 'Max drawdown %', 'Account blown?',
                    'Profit / max drawdown'], height=42); r += 1
        for (tf, prof), rows in C.items():
            for x in rows:
                q = x['r']; fl = BAD if q['ruined'] else None
                vals = [tf, PNAME[prof], 'Rule C' if x['mode'].startswith('Rule C:') else 'Rule C split', x['split'], '$%d' % x['cap'], 'Clamp' if x['act'] == 'clamp' else 'Start again',
                        round(q['net'], 2), round(q['maxDD'], 2), round(q['maxDDpct'], 4), 'YES' if q['ruined'] else 'no', '=IF(H{0}=0,"",G{0}/H{0})']
                for j, v in enumerate(vals, 1):
                    isf = isinstance(v, str) and v.startswith('=')
                    put(ws, r, j, v.format(r) if isf else v, '#,##0;(#,##0);-' if j in (7, 8) else '0.0%' if j == 9 else '0.00' if j == 11 else None, False,
                        BLUE if 7 <= j <= 9 else None, fl)
                r += 1
            r += 1
    fs = S + '/gold/rulec_stress.pkl'
    if os.path.exists(fs):
        T = pickle.load(open(fs, 'rb'))
        r += 1; put(ws, r, 1, 'PART 3 - stress test on the best settings: the same runs with 0.05 R and 0.10 R of extra cost on every trade (slippage, a weaker live edge)',
                     bold=True, fill=SUB); ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=12); r += 1
        hdr(ws, r, ['Timeframe', 'Rule', 'Split', 'Hard cap / action', 'Net ($), as tested', 'Max DD %, as tested', 'Net ($), +0.05 R cost', 'Max DD %, +0.05 R', 'Net ($), +0.10 R cost',
                    'Max DD %, +0.10 R'], height=42); r += 1
        keys = sorted({k[2:] for k in T}, key=lambda k: (k[0] != 'Off', k[0], k[1], k[3]))
        for tf in TFS:
            for mode, split, cap, act in keys:
                if (tf, 0.0, mode, split, cap, act) not in T: continue
                vals = [tf, 'Off (fixed $50)' if mode == 'Off' else ('Rule C' if mode.startswith('Rule C:') else 'Rule C split'), '-' if mode == 'Off' else split,
                        '-' if not cap else '$%d %s' % (cap, 'Clamp' if act == 'clamp' else 'Start again')]
                for extra in (0.0, 0.05, 0.10):
                    q = T[(tf, extra, mode, split, cap, act)]
                    vals += ['BLOWN' if q['ruined'] else round(q['net'], 2), round(q['maxDDpct'], 4)]
                for j, v in enumerate(vals, 1):
                    put(ws, r, j, v, '#,##0;(#,##0);-' if j in (5, 7, 9) else '0.0%' if j in (6, 8, 10) else None, mode == 'Off', BLUE if j >= 5 else None,
                        TOT if mode == 'Off' else (BAD if v == 'BLOWN' else None))
                r += 1
            r += 1
    widths(ws, [9, 30, 14, 6, 9, 11, 11, 9, 9, 9, 16, 9, 9, 9, 10, 11, 11]); ws.freeze_panes = 'D11'
# ---------------------------------------------------------------- robustness (neighbourhood of each pick)
def robust_sheet(wb, FIN):
    ws = wb.create_sheet('Robustness')
    title(ws, 'How fragile are the recommended settings? Every single setting changed once more on top of them',
          'A setting that only works at one exact value is probably luck. Shown: how many one-step changes stay profitable after spread in both halves, and the changes that hurt most.')
    ws.merge_cells('A2:J2'); r = 4
    cols = ['Timeframe', 'Settings', 'One-step changes tested', 'Still profitable in 2021-23 and 2024-26', 'Share', 'Median R per trade after spread (all changes)',
            'R per trade of the pick itself', 'Most damaging change', 'R per trade after it', '2nd most damaging', 'R per trade after it', '3rd most damaging', 'R per trade after it']
    hdr(ws, r, cols, height=54); r += 1
    for tf in TFS:
        if tf not in FIN or 'neigh' not in FIN[tf]: continue
        for prof in ('profit', 'streak'):
            if prof not in FIN[tf]['neigh']: continue
            rows = []
            for mn, c, tr in FIN[tf]['neigh'][prof]:
                if not tr: continue
                m = tr if isinstance(tr, dict) else greedy.metrics(greedy.pack(tr), {}); rows.append((mn, m))
            if not rows: continue
            both = sum(1 for _, m in rows if m['A']['E'] > 0 and m['B']['E'] > 0)
            med = statistics.median(m['all']['E'] for _, m in rows)
            own = [x[1] - x[4] / x[5] for x in FIN[tf]['trades'][prof]]
            worst = sorted(rows, key=lambda x: x[1]['all']['E'])[:3]
            vals = [tf, PNAME[prof], len(rows), both, '=IF(C{0}=0,"",D{0}/C{0})', round(med, 4), round(sum(own) / len(own), 4) if own else 0]
            for mn, m in worst: vals += [mn, round(m['all']['E'], 4)]
            for j, v in enumerate(vals, 1):
                isf = isinstance(v, str) and v.startswith('=')
                put(ws, r, j, v.format(r) if isf else v, '0.0%' if j == 5 else '0.000' if j in (6, 7, 9, 11, 13) else None, False,
                    None if (isf or j in (1, 2, 8, 10, 12)) else BLUE, wrap=j in (8, 10, 12))
            r += 1
        r += 1
    widths(ws, [9, 30, 9, 11, 7, 11, 9, 30, 9, 30, 9, 30, 9]); ws.freeze_panes = 'C5'
# ---------------------------------------------------------------- hour of day
def hours_sheet(wb, FIN):
    ws = wb.create_sheet('Hours')
    title(ws, 'When the default settings trade (IST hour of the fill) and what the spread costs then',
          'With max 5 fills a day, M1 uses all 5 fills every day in the quiet hours just after the broker midnight (VWAP reset) - where the spread is a big part of the small stop.')
    ws.merge_cells('A2:Q2'); r = 4
    tfs = [tf for tf in TFS if tf in FIN]
    put(ws, r, 1, 'IST hour', bold=True, fill=SUB)
    for k, tf in enumerate(tfs):
        c0 = 2 + 4 * k; put(ws, r, c0, tf, bold=True, fill=SUB); ws.merge_cells(start_row=r, start_column=c0, end_row=r, end_column=c0 + 3)
    r += 1; hdr(ws, r, ['IST hour'] + ['Trades', 'Share', 'Spread cost (R per trade)', 'R per trade after spread'] * len(tfs), height=42); r += 1
    stats = {}
    for tf in tfs:
        tr = FIN[tf]['trades']['default']; by = collections.defaultdict(list)
        for x in tr: by[x[0].hour].append(x)
        stats[tf] = (len(tr), by)
    first = r
    for h in range(24):
        put(ws, r, 1, '%02d:00' % h, bold=True)
        for k, tf in enumerate(tfs):
            n, by = stats[tf]; q = by.get(h, []); c0 = 2 + 4 * k
            put(ws, r, c0, len(q), None, False, BLUE)
            put(ws, r, c0 + 1, '=IF(SUM(%s$%d:%s$%d)=0,"",%s%d/SUM(%s$%d:%s$%d))' % (CL(c0), first, CL(c0), first + 23, CL(c0), r, CL(c0), first, CL(c0), first + 23), '0.0%')
            put(ws, r, c0 + 2, round(sum(x[4] / x[5] for x in q) / len(q), 4) if q else '', '0.00', False, BLUE)
            put(ws, r, c0 + 3, round(sum(x[1] - x[4] / x[5] for x in q) / len(q), 4) if q else '', '0.000', False, BLUE,
                (GOOD if q and sum(x[1] - x[4] / x[5] for x in q) > 0 else BAD) if q else None)
        r += 1
    put(ws, r, 1, 'Total', bold=True, fill=TOT)
    for k, tf in enumerate(tfs):
        c0 = 2 + 4 * k; put(ws, r, c0, '=SUM(%s%d:%s%d)' % (CL(c0), first, CL(c0), r - 1), None, True, None, TOT)
        for j in (1, 2, 3): put(ws, r, c0 + j, '', fill=TOT)
    widths(ws, [9] + [8, 7, 9, 9] * len(tfs)); ws.freeze_panes = 'B6'
# ---------------------------------------------------------------- streak math
def math_sheet(wb, FIN):
    ws = wb.create_sheet('Streak math')
    title(ws, 'Why the losing streak depends on the win rate (and only a little on the filters)',
          'Expected longest losing streak in N independent trades with win rate p: ln(N x p) / ln(1/(1-p)) + 0.5772 / ln(1/(1-p)) - 0.5 (Schilling). Edit the blue cells.')
    ws.merge_cells('A2:J2')
    put(ws, 4, 1, 'Win rate p ->', bold=True); ns = [100, 250, 500, 1000, 2000, 3500, 7000]; ps = [0.30, 0.33, 0.36, 0.40, 0.45, 0.50, 0.55, 0.60, 0.65, 0.70]
    put(ws, 5, 1, 'Trades N', bold=True, fill=SUB)
    for j, p in enumerate(ps, 2): put(ws, 5, j, p, '0%', True, BLUE, SUB)
    for i, n in enumerate(ns, 6):
        put(ws, i, 1, n, '#,##0', True, BLUE)
        for j in range(2, 2 + len(ps)):
            L = CL(j); put(ws, i, j, '=LN($A{0}*{1}$5)/LN(1/(1-{1}$5))+0.5772/LN(1/(1-{1}$5))-0.5'.format(i, L), '0.0')
    r = 6 + len(ns) + 1
    put(ws, r, 1, 'Expected longest WIN streak (same formula with p and 1-p swapped)', bold=True); r += 1
    put(ws, r, 1, 'Trades N', bold=True, fill=SUB)
    for j in range(2, 2 + len(ps)): put(ws, r, j, '=%s5' % CL(j), '0%', True, None, SUB)
    hr = r; r += 1
    for i, n in enumerate(ns):
        put(ws, r, 1, '=A%d' % (6 + i), '#,##0', True)
        for j in range(2, 2 + len(ps)):
            L = CL(j); put(ws, r, j, '=LN($A{0}*(1-{1}${2}))/LN(1/{1}${2})+0.5772/LN(1/{1}${2})-0.5'.format(r, L, hr), '0.0')
        r += 1
    r += 1
    notes = ['How to read it: the default settings win about 33% at TP4 (2R). Over ~900 H1 trades the table expects a longest losing streak of ~15 - the backtest had 15.',
             'Exiting at TP1 (0.6R) wins about 60%: the expected longest losing streak drops to ~7, as in the backtest - but 0.6R wins need > 62.5% plus spread to make money.',
             'Losses in this backtest do not cluster: the chance of a loss after 1, 2, 3... losses in a row stays at the overall loss rate, so pausing after losses does not shorten streaks.',
             'So the ways to shorter losing streaks are: a higher win rate (nearer target, breakeven move, partial close at TP1) or fewer trades - not pause / stop-for-the-day rules.']
    for t in notes: ws.cell(r, 1, t).font = font(size=9); r += 1
    r += 1; put(ws, r, 1, 'Do losses cluster? Chance that the next trade loses after k losses in a row (default settings, exit TP4, 2021-2026)', bold=True, fill=SUB)
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=10); r += 1
    hdr(ws, r, ['Timeframe', 'Trades', 'Loss rate (all trades)', 'After 1 loss', 'After 2 losses', 'After 3 losses', 'After 4 losses', 'After 5+ losses', 'Cases after 5+'], height=30); r += 1
    for tf in TFS:
        if tf not in FIN: continue
        res = [x[1] for x in sorted(FIN[tf]['trades']['default'], key=lambda x: x[0])]
        cnt = collections.defaultdict(lambda: [0, 0]); run = 0
        for v in res:
            k = min(run, 5)
            if k: cnt[k][0] += 1; cnt[k][1] += v < 0
            run = run + 1 if v < 0 else 0
        vals = [tf, len(res), sum(1 for v in res if v < 0) / len(res)] + [(cnt[k][1] / cnt[k][0]) if cnt[k][0] else '' for k in (1, 2, 3, 4, 5)] + [cnt[5][0]]
        for j, v in enumerate(vals, 1): put(ws, r, j, v, '0.0%' if 3 <= j <= 8 else None, False, BLUE if j > 1 else None)
        r += 1
    ws.cell(r, 1, 'If losses clustered, the chance after several losses would be clearly higher than the overall rate. It is not, so a streak is just bad luck in a row.').font = font(size=9, italic=True)
    widths(ws, [12] + [9] * 10)
# ---------------------------------------------------------------- best settings
def best_sheet(wb, FIN, refs, rbest):
    ws = wb.create_sheet('Best settings', 1)
    title(ws, 'Recommended settings per timeframe (FxPro MT5 GOLD, 2021-2026)', None)
    ws['A2'] = ('Picked from every configuration tested (one-at-a-time sweep, combinations, step-by-step search; thousands per timeframe). Rule: profitable after spread in BOTH '
                '2021-23 and 2024-26, at least 4 of 6 years positive, enough trades; ranked by the weaker half. Because both halves were used to pick, expect live results to be '
                'weaker than shown (the step-by-step search chosen on 2021-23 alone shows how much tuning gains fade out of sample).')
    ws['A2'].font = font(italic=True); ws.merge_cells('A2:H2'); ws['A2'].alignment = Alignment(wrap_text=True, vertical='top'); ws.row_dimensions[2].height = 54
    ws['A3'] = ('Recovery rules: across all timeframes fixed risk (Off) gave the best profit per unit of drawdown on average - the line per setting below shows the single best rule '
                'for that setting (never risking more than $500 on one trade, never falling more than 40%). See the Recovery rules sheet. With the best (profit) settings, Rule C split at split 3 '
                'with the $500 cap made 5.8x (H1) and 9.2x (M15) the profit of fixed risk with about the same loss per $ staked, but its drawdowns were 20% and 33% instead of 3% and 9%, '
                'and without the cap it blew the account on M15 - see Rule C split test.')
    ws['A3'].font = font(italic=True); ws.merge_cells('A3:H3'); ws['A3'].alignment = Alignment(wrap_text=True, vertical='top'); ws.row_dimensions[3].height = 40
    r = 5; tot = {(tf, prof): (sheet, t) for tf, prof, sheet, (_, t) in refs}
    for tf in TFS:
        if tf not in FIN: continue
        put(ws, r, 1, tf, bold=True, fill=TOT); ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=8); r += 1
        for prof in ('profit', 'streak'):
            cfg = FIN[tf]['cfg'][prof]; put(ws, r, 1, PNAME[prof], bold=True, fill=SUB); ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=8); r += 1
            if FIN[tf].get('passed', {}).get(prof) is False:
                put(ws, r, 1, 'WARNING: no setting tested on %s was profitable after spread in both 2021-23 and 2024-26. These are the least-bad settings found - do not trade them expecting a profit.' % tf,
                    bold=True, fill=BAD, wrap=True)
                ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=8); ws.row_dimensions[r].height = 26; r += 1
            hdr(ws, r, ['Group', 'Indicator / strategy input', 'Set to', 'Default', 'MT5 EA input'], height=18); r += 1
            rows = labels.rows(cfg) or [('-', 'no change: default settings', '-', '-', '-')]
            for g in rows:
                for j, v in enumerate(g, 1): put(ws, r, j, v, wrap=(j in (2, 5)))
                r += 1
            sheet, t = tot[(tf, prof)]; q = "'%s'!" % sheet
            stats = [('Trades 2021-2026', '=%sC%d' % (q, t), '#,##0'), ('Trades per day', '=%sC%d/%d' % (q, t, DAYS_N), '0.00'),
                     ('Win % (closed at final target)', '=%sP%d' % (q, t), '0.0%'), ('Profitable trades %', '=%sQ%d' % (q, t), '0.0%'),
                     ('TP1 / TP2 / TP3 / TP4 reached', '=TEXT(%sR%d,"0%%")&" / "&TEXT(%sS%d,"0%%")&" / "&TEXT(%sT%d,"0%%")&" / "&TEXT(%sU%d,"0%%")' % (q, t, q, t, q, t, q, t), None),
                     ('R per trade after spread', '=%sW%d' % (q, t), '0.000'), ('Total R after spread', '=%sM%d' % (q, t), '0.0'),
                     ('Longest win / losing streak', '=%sN%d&" / "&%sO%d' % (q, t, q, t), None)]
            for k, (lab, f, nf) in enumerate(stats):
                rr = r + k // 2; cc = 1 + 4 * (k % 2)
                put(ws, rr, cc, lab, bold=True); c = put(ws, rr, cc + 1, f, nf, False, GREEN)
            r += 4
            rb = [x for x in rbest if x[0] == tf and x[1] == prof]
            if rb:
                _, _, b, off = rb[0]
                put(ws, r, 1, 'Best recovery rule', bold=True)
                if b['mode'] == 'Off':
                    txt = 'Off (fixed $50 risk) gives the best profit for the drawdown: net $%s, max drawdown $%s (%.0f%%) on $10,000 - no recovery rule beat it.' % (
                        format(round(b['net']), ','), format(round(b['maxDD']), ','), 100 * b['maxDDpct'])
                else:
                    txt = '%s%s: net $%s, max drawdown $%s (%.0f%%), largest single risk $%s - vs Off: net $%s, max drawdown $%s (%.0f%%).' % (
                        b['mode'], ' + $500 cap' if b['cap'] else '', format(round(b['net']), ','), format(round(b['maxDD']), ','), 100 * b['maxDDpct'],
                        format(round(b['maxRisk']), ','), format(round(off['net']), ','), format(round(off['maxDD']), ','), 100 * off['maxDDpct'])
                put(ws, r, 2, txt, wrap=True)
                ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=8); ws.row_dimensions[r].height = 24; r += 1
            r += 1
        r += 1
    widths(ws, [30, 46, 26, 12, 30, 12, 10, 10])
# ---------------------------------------------------------------- read me
def readme(wb, FIN):
    ws = wb.active; ws.title = 'Read me'
    npool = {}
    for tf in TFS:
        fn = S + '/gold/greedy_pool_%s.pkl' % tf
        if os.path.exists(fn): npool[tf] = len(pickle.load(open(fn, 'rb')))
    L = ['Auction Retest indicator on FxPro MT5 GOLD, 4 Jan 2021 - 2 Oct 2026 (your 1-minute data): profitable settings per goal, accuracy, TP1-TP4, streaks, recovery rules', '',
         'HOW TO USE THIS FILE:',
         '  1. Strategy guide: pick your goal (most profit, shortest losing streak, longest winning streak, highest win rate, most trades, every year positive, smoothest, simplest,',
         '     one session only) and your timeframe. Click the strategy to jump to its column.',
         '  2. <TF> strategies (M1, M5, M15, H1): one column per strategy. Every indicator input is listed with its exact name; yellow = change it to that value, grey = keep the default.',
         '     Below the settings: what to expect - trades, win %, TP1-TP4 reached, R per trade in each half and each year, longest winning and losing streaks, deepest fall, $ results.',
         '  3. All profitable settings: every setting that passed the profit rule, with a filter on each column (e.g. longest losing streak <= 6).',
         '  4. Recovery rules and Rule C split test: money management. The rest of the sheets are the research detail.', '',
         'DATA: 2,037,975 one-minute bars from your MT5 export; 441 fake 00:00 "daily summary" bars removed. M5 / M15 / H1 were built from M1 on the broker clock (EET, UTC+2 / UTC+3 in summer).',
         'ENGINE: the same signal engine as the TradingView indicator / strategy and the MT5 EA (Python twin, checked trade-for-trade against the EA). VWAP and sessions reset at broker midnight; VWAP weights = tick volume.',
         'COSTS: the spread recorded in the MT5 data at each fill bar, as a share of the stop distance ("after spread"). On M1 / M5 the stop is small, so the spread eats a large part of every trade.',
         'DEFINITIONS: R = result in units of the stop distance. TPk hit = price reached that target after the fill bar and before the stop. Win = closed at the final target.',
         '  Profitable = result above 0 (with a partial close at TP1 a trade that reached TP1 is profitable). A breakeven exit is neither a win nor a loss and ends both streaks.',
         '  Longest win / losing streak = most consecutive profitable / losing trades. Months and years use the IST date of the fill.',
         'HOW THE SETTINGS WERE SEARCHED (choose on 2021-23, check on 2024-26):',
         '  1. Every setting changed alone, all others default: 146 versions per timeframe (sheet Settings sweep).',
         '  2. Every combination of 1-4 of the 8 settings that helped most alone: ~160 per timeframe for H1, M15, M5 (sheet Combinations, part 1).',
         '  3. Exit choices (each TP, breakeven moves, partial) for the best candidates on H1 and M15 (sheet Exits).',
         '  4. Step-by-step search: each step tries every single change (~170: all settings incl. extra values, TP levels, stop moves, resets; M5 109 and M1 61 to keep run time sane)',
         '     on top of the current settings and keeps the one that helps 2021-23 most (sheet Combinations, part 2).',
         '  5. Every version was also scored with 30 / 50 / 70% closed at TP1 (partial close: strategy and EA only).',
         '  Settings scored in total: ' + ', '.join('%s %s' % (tf, format(n, ',')) for tf, n in npool.items()) + '. Literally every combination is impossible (~60 settings: more than 10^40 combinations).',
         'PICK RULE: profitable after spread in BOTH 2021-23 and 2024-26, at least 4 of 6 years positive, enough trades; ranked by the weaker half; then the version with the fewest changes',
         '  within 10% of the best. Low-streak pick = the shortest longest-losing-streak among settings that pass the same rule and keep a reasonable profit.',
         '  Because both halves were used to pick, expect live results to be weaker than these tables. Sheet Robustness shows how the picks behave when any one setting is changed again.',
         'COLOURS: blue = backtest results, black = formulas, green = links to another sheet. Partial close and recovery rules exist in the strategy and the EA only (not in the indicator).', '',
         'SHEETS: Strategy guide | M1 / M5 / M15 / H1 strategies | All profitable settings | Recovery rules | Rule C split test | Summary | <TF> default / best / low-streak (year x month;',
         '  best = strategy S1 "Most profit", low-streak = the "Shortest losing streak" strategy) | TP1-TP4 | Settings sweep | Combinations | Exits | Robustness | Hours | Streak math']
    for i, t in enumerate(L, 1): ws.cell(i, 1, t).font = Font(name=F, bold=(i == 1), size=11 if i == 1 else 10)
    ws.column_dimensions['A'].width = 190
if __name__ == '__main__':
    FIN = {}
    for tf in TFS:
        fn = S + '/gold/final_%s.pkl' % tf
        if os.path.exists(fn):
            d = pickle.load(open(fn, 'rb'))
            if all(k in d for k in ('cfg', 'trades', 'exits', 'recovery')): FIN[tf] = d
    wb = Workbook(); readme(wb, FIN); refs = []
    for tf in TFS:
        if tf not in FIN: continue
        for prof, suffix in PROF:
            sheet = '%s %s' % (tf, suffix)
            yrows, t = ym_sheet(wb, sheet, FIN[tf]['trades'][prof], '%s - %s - accuracy per year and month' % (tf, PNAME[prof]), FIN[tf]['cfg'][prof])
            refs.append((tf, prof, sheet, (yrows, t)))
    summary(wb, refs); tp_sheet(wb, FIN); sweep_sheet(wb); combo_sheet(wb); exits_sheet(wb); rbest = recovery_sheet(wb, FIN); rulec_sheet(wb); robust_sheet(wb, FIN); hours_sheet(wb, FIN); math_sheet(wb, FIN)
    import strategy_sheets as SS
    CAT = SS.load()
    for tf in CAT:                       # the month sheets 'best' / 'low-streak' are the catalogue's Most profit / Shortest losing streak strategies
        names = {n: s['cfg'] for s in CAT[tf]['strategies'] for n, _ in s['scen']}
        assert names.get('Most profit') == FIN[tf]['cfg']['profit'] and names.get('Shortest losing streak') == FIN[tf]['cfg']['streak'], tf
    links = {tf: SS.tf_sheet(wb, tf, CAT[tf]) for tf in TFS if tf in CAT}
    SS.guide(wb, CAT, links); SS.allprof(wb, CAT)
    order = (['Read me', 'Strategy guide'] + ['%s strategies' % tf for tf in TFS] + ['All profitable settings', 'Recovery rules', 'Rule C split test', 'Summary'] +
             [s for _, _, s, _ in refs] + ['TP1-TP4', 'Settings sweep', 'Combinations', 'Exits', 'Robustness', 'Hours', 'Streak math'])
    wb._sheets = [wb[n] for n in order if n in wb.sheetnames]
    out = sys.argv[1] if len(sys.argv) > 1 else S + '/gold/Gold_2021_2026_Auction_Retest.xlsx'
    wb.save(out); print('saved', out, 'timeframes', list(FIN))
